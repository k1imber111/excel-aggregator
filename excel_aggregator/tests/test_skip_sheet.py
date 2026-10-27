"""Команда «п» на шаге 3: пропустить лист — он не попадает в итог.

Приложение гоняется в памяти с заранее заданными ответами. Книга из трёх
листов с общими названиями объектов A, B, C; на каждом — своя характеристика.
"""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest
from rich.console import Console

from excel_aggregator.app import AggregatorApp

_OBJECTS = [("A", 1), ("B", 2), ("C", 3)]


def _book(path: Path, sheets: dict[str, list[str]]) -> Path:
    """Книга: лист -> шапка «Название» + перечисленные характеристики."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, titles in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(["Название", *titles])
        for key, n in _OBJECTS:
            ws.append([key, *[None if t == "Пусто" else n * 10 + i for i, t in enumerate(titles)]])
    wb.save(path)
    return path


@pytest.fixture
def three_sheets(tmp_path: Path) -> Path:
    return _book(tmp_path / "книга.xlsx", {
        "Нужный": ["Высота", "Ширина"],
        "Лишний": ["Мусор"],
        "Ещё нужный": ["Вес"],
    })


def _drive(answers: list[str], monkeypatch) -> tuple[AggregatorApp, str]:
    monkeypatch.setenv("AGGREGATOR_SCRIPTED", "1")
    app = AggregatorApp()
    buf = io.StringIO()
    app.ui.console = Console(file=buf, width=200)
    queue = list(answers)
    app.ui.console.input = lambda prompt="": queue.pop(0)
    app.run()
    assert not queue, f"Остались непрочитанные ответы: {queue}"
    return app, buf.getvalue()


def _result(path: Path) -> tuple[list, list]:
    """Нижняя строка шапки и строки данных итогового файла."""
    ws = openpyxl.load_workbook(path).active
    leaf = int(ws.auto_filter.ref.split(":")[0][1:])
    headers = [ws.cell(row=leaf, column=c).value for c in range(1, ws.max_column + 1)]
    headers[0] = ws["A2"].value     # заголовок первого столбца объединён по вертикали
    return headers, list(ws.iter_rows(min_row=leaf + 1, values_only=True))


def _flow(path: Path, step3: list[str], out: Path) -> list[str]:
    """Файл -> проверка -> ответы шага 3 -> сохранение без открытия -> выход."""
    return ["1", str(path), "", "", *step3, str(out), "н", "н", "3"]


def test_skip_middle_sheet(three_sheets, tmp_path, monkeypatch):
    out = tmp_path / "итог"
    app, _ = _drive(_flow(three_sheets, ["", "п", ""], out), monkeypatch)
    headers, rows = _result(tmp_path / "итог.xlsx")
    assert headers == ["Название", "Высота", "Ширина", "Вес"]
    assert rows == [("A", 10, 11, 10), ("B", 20, 21, 20), ("C", 30, 31, 30)]
    assert [t.enabled for t in app.tables] == [True, False, True]


def test_skip_first_sheet(three_sheets, tmp_path, monkeypatch):
    out = tmp_path / "итог"
    _drive(_flow(three_sheets, ["п", "", ""], out), monkeypatch)
    headers, rows = _result(tmp_path / "итог.xlsx")
    assert headers == ["Название", "Мусор", "Вес"]
    assert rows == [("A", 10, 10), ("B", 20, 20), ("C", 30, 30)]


def test_skip_last_sheet(three_sheets, tmp_path, monkeypatch):
    out = tmp_path / "итог"
    _drive(_flow(three_sheets, ["", "", "п"], out), monkeypatch)
    headers, _ = _result(tmp_path / "итог.xlsx")
    assert headers == ["Название", "Высота", "Ширина", "Мусор"]


@pytest.mark.parametrize("word", ["п", "П", "далее", "Пропустить", "g"])
def test_skip_command_spellings(three_sheets, tmp_path, monkeypatch, word):
    out = tmp_path / "итог"
    _drive(_flow(three_sheets, ["", word, ""], out), monkeypatch)
    headers, _ = _result(tmp_path / "итог.xlsx")
    assert "Мусор" not in headers


def test_last_remaining_sheet_cannot_be_skipped(three_sheets, tmp_path, monkeypatch):
    out = tmp_path / "итог"
    # третий «п» отклонён: иначе собирать было бы нечего; затем выбран столбец «Вес»
    _, text = _drive(_flow(three_sheets, ["п", "п", "п", ""], out), monkeypatch)
    assert "единственный" in text
    headers, _ = _result(tmp_path / "итог.xlsx")
    assert headers == ["Название", "Вес"]


def test_back_to_skipped_sheet_and_reenable(three_sheets, tmp_path, monkeypatch):
    out = tmp_path / "итог"
    # лист 1 — «п»; на листе 2 «н» возвращает к листу 1; выбираем только «Высота»
    app, _ = _drive(_flow(three_sheets, ["п", "н", "1", "", ""], out), monkeypatch)
    headers, _ = _result(tmp_path / "итог.xlsx")
    assert headers == ["Название", "Высота", "Мусор", "Вес"]
    assert all(t.enabled for t in app.tables)


def test_skip_then_build_takes_all_for_the_rest(three_sheets, tmp_path, monkeypatch):
    out = tmp_path / "итог"
    # лист 1 пропущен, на листе 2 команда «с»: остальные берутся целиком
    _drive(_flow(three_sheets, ["п", "с"], out), monkeypatch)
    headers, _ = _result(tmp_path / "итог.xlsx")
    assert headers == ["Название", "Мусор", "Вес"]


def test_skipped_sheet_is_marked_on_review(three_sheets, tmp_path, monkeypatch):
    out = tmp_path / "итог"
    # «п» на листе 1, «н» к нему же, «н» ещё раз — на шаг 2; потом снова шаг 3
    answers = _flow(three_sheets, ["п", "н", "н", "", "", ""], out)
    _, text = _drive(answers, monkeypatch)
    assert "не используется" in text
    headers, _ = _result(tmp_path / "итог.xlsx")
    assert headers == ["Название", "Мусор", "Вес"]


def test_skip_sheets_in_two_files(tmp_path, monkeypatch):
    first = _book(tmp_path / "первый.xlsx", {"a1": ["Высота"], "a2": ["Лишнее1"]})
    second = _book(tmp_path / "второй.xlsx", {"b1": ["Лишнее2"], "b2": ["Вес"]})
    out = tmp_path / "итог"
    answers = ["1", str(first), str(second), "", "", "", "п", "п", "", str(out), "н", "н", "3"]
    _drive(answers, monkeypatch)
    headers, _ = _result(tmp_path / "итог.xlsx")
    assert headers == ["Название", "Высота", "Вес"]


def test_back_skips_sheet_without_columns(tmp_path, monkeypatch):
    # Лист 2 без столбцов характеристик пропускается сам. «н» на листе 3 должна
    # вести к листу 1, а не возвращать на лист 3.
    path = _book(tmp_path / "книга.xlsx", {
        "Первый": ["Высота", "Ширина"],
        "Пустой": ["Пусто"],
        "Третий": ["Вес"],
    })
    out = tmp_path / "итог"
    app, _ = _drive(_flow(path, ["", "н", "1", ""], out), monkeypatch)
    assert app.tables[1].chosen == []
    headers, _ = _result(tmp_path / "итог.xlsx")
    assert headers == ["Название", "Высота", "Вес"]


def test_back_from_first_sheet_goes_to_review(three_sheets, tmp_path, monkeypatch):
    out = tmp_path / "итог"
    # «н» на первом листе — возврат на проверку (Enter), затем шаг 3 заново
    _, text = _drive(_flow(three_sheets, ["н", "", "", "", ""], out), monkeypatch)
    assert text.count("Шаг 2 из 3. Проверка") == 2
