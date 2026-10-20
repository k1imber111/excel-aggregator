"""Красная команда, волна 2: новые правила matcher / keys / exporter.

Правила: нумерация 1..N не соединяет таблицы с текстовыми названиями,
заголовок-измерение штрафуется, _SUMMARY_RE различает итоги и объекты,
exporter пишет большие целые текстом, длительности как [HH]:MM:SS и чистит текст.
"""

from __future__ import annotations

import math
import random
from datetime import date, datetime, timedelta

import openpyxl
import pytest

from excel_aggregator.core import Orientation, SheetGrid
from excel_aggregator.core.errors import AppError
from excel_aggregator.core.exporter import export
from excel_aggregator.core.keys import extract_keys
from test_join_adversarial import objects, pipeline, val, write_xlsx


# ---------------------------------------------------------------- matcher


@pytest.mark.parametrize(
    "title", ["Площадка", "Массив", "Высотное сооружение", "Угольный разрез", "Весы"]
)
def test_names_column_with_measure_like_title_still_joins_by_names(tmp_path, title):
    """Столбец названий с заголовком «Площадка»/«Массив»… (регэксп видит «площад», «масс»)
    получает штраф 0.3, проигрывает столбцу уникальных целых и не соединяется:
    те же 8 объектов дают 16 строк."""
    names = objects(8, "Объект")
    rnd = random.Random(1)

    def table(src):
        regs = rnd.sample(range(1000, 9999), len(names))   # свои рег. номера в каждом файле
        return [[title, "Рег. номер", "П"]] + [
            [n, regs[i], val(src, i + 1, 1)] for i, n in enumerate(names)
        ]
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": table(1)})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": table(2)})
    _, tables, _, _, rows = pipeline([pa, pb])
    assert [t.key_col for t in tables] == [0, 0], [t.key_col for t in tables]
    assert len(rows) == 8, f"объекты не соединены: {len(rows)} строк вместо 8"


@pytest.mark.parametrize("second", ["address", "date"])
def test_real_row_number_key_joins_when_tables_have_other_unique_text(tmp_path, second):
    """Таблицы связаны только по № (1..N); у каждой есть свой уникальный текстовый
    столбец («Адрес»/«Комментарий», даты осмотра). is_text_names считает его
    «названиями» и обнуляет пару №—№: 10 общих объектов дают 20 строк."""
    def table(src):
        if second == "address":
            extra = [f"ул. {'Аб'[src - 1]}{i}" for i in range(1, 11)]
            head = ["№", "Адрес" if src == 1 else "Комментарий", "П"]
        else:
            extra = [datetime(2024 + src, 1, i) for i in range(1, 11)]
            head = ["№", "Дата осмотра", "П"]
        return [head] + [[i, extra[i - 1], val(src, i, 1)] for i in range(1, 11)]
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": table(1)})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": table(2)})
    _, tables, _, _, rows = pipeline([pa, pb])
    assert [t.key_col for t in tables] == [0, 0], [t.key_col for t in tables]
    assert len(rows) == 10, len(rows)


def test_numbering_joins_when_only_one_table_has_names_but_no_cross_row(tmp_path):
    """Контроль: таблица с названиями и таблица только с № не должны сливать
    чужие строки — либо соединены по № (одинаковые ключи), либо раздельно.
    Ячейка каждого объекта — из его же строки источника."""
    na = objects(10)
    ta = [["№", "Название", "П"]] + [[i, n, val(1, i, 1)] for i, n in enumerate(na, 1)]
    tb = [["№", "Q"]] + [[i, val(2, i, 1)] for i in range(1, 11)]
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": tb})
    _, _, _, _, rows = pipeline([pa, pb])
    for r in rows:
        vals = [v for v in r[1:] if isinstance(v, int) and v >= 100000]
        idx = {v % 100000 // 10 for v in vals}
        assert len(idx) <= 1, f"в строке {r[0]!r} данные разных объектов: {r}"


def test_four_sources_numbers_and_names_do_not_cross_cells(tmp_path):
    """4 источника: два по названиям (разный порядок и нумерация 1..N), два
    только по №. Данные одного объекта не попадают в чужую строку."""
    common = objects(6)
    a = common + [f"Опора {i}" for i in range(1, 7)]
    c = list(reversed(common)) + [f"Антенна {i}" for i in range(1, 7)]
    ids = {n: i for i, n in enumerate(dict.fromkeys(a + c), 1)}

    def named(src, ns):
        return [["№", "Название", "П"]] + [
            [i, n, val(src, ids[n], 1)] for i, n in enumerate(ns, 1)
        ]
    paths = [
        write_xlsx(tmp_path / "a.xlsx", {"Л": named(1, a)}),
        write_xlsx(tmp_path / "c.xlsx", {"Л": named(2, c)}),
        write_xlsx(tmp_path / "d.xlsx", {"Л": [["№", "К"]] + [[i, 300000 + i] for i in range(1, 13)]}),
        write_xlsx(tmp_path / "e.xlsx", {"Л": [["№", "Л"]] + [[i, 400000 + i] for i in range(1, 13)]}),
    ]
    _, _, _, headers, rows = pipeline(paths)
    for r in rows:
        key = r[0]
        if key in ids:
            for v in r[1:]:
                if isinstance(v, int) and 100000 <= v < 300000:
                    assert v % 100000 // 10 == ids[key], f"{key}: чужое значение {v}"


# ------------------------------------------------------------------- keys


def _keys(labels: list[str]) -> list[str]:
    rows = [["Мачта 1", 5], ["Мачта 2", 6]] + [[x, 100] for x in labels]
    return extract_keys(SheetGrid("s", rows), 0, 0, Orientation.VERTICAL).keys


@pytest.mark.parametrize(
    "label",
    [
        "Итого за год",
        "Итого (руб.)",
        "Итог",
        "Итог:",
        "Всего за период",
        "Итого - по району",
        "Всего в районе",
        "Итого на конец года",
        "Total for 2024",
        "Subtotal",
    ],
)
def test_real_total_label_with_numbers_is_not_an_object(label):
    """Итоговая строка с суммами в соседнем столбце не должна становиться объектом."""
    keys = _keys([label])
    assert label not in keys, f"итог {label!r} стал объектом: {keys}"


@pytest.mark.parametrize("label", ["Total, Inc", "Всего, ООО"])
def test_object_with_punctuation_after_total_word_is_kept(label):
    """Объект «Total, Inc» с данными: правило `\\s*[:;.,]` считает его итогом."""
    assert label in _keys([label]), f"объект {label!r} потерян как подпись"


@pytest.mark.parametrize(
    "label", ["Итого", "ИТОГО:", "Итого 5", "Всего по району", "Примечание: ок", "Примечания"]
)
def test_known_summary_labels_stay_excluded(label):
    assert label not in _keys([label])


@pytest.mark.parametrize("label", ["Всего Мачта", "Total Tower", "Итого-Север", "total_station 5"])
def test_object_names_with_total_words_stay(label):
    assert label in _keys([label])


# --------------------------------------------------------------- exporter


def _read(path):
    return openpyxl.load_workbook(path).active


def test_nan_and_inf_do_not_corrupt_or_crash(tmp_path):
    out = tmp_path / "o.xlsx"
    for bad in (float("nan"), float("inf"), float("-inf")):
        try:
            export(out, ["Код", "v"], [["a", bad], ["b", 1.5]])
        except AppError:
            continue
        v = _read(out)["B2"].value
        assert v is None or isinstance(v, str) or (isinstance(v, float) and math.isfinite(v)), v


@pytest.mark.parametrize(
    "value",
    [datetime(1850, 1, 1), date(1, 1, 1), datetime(9999, 12, 31, 23, 59), datetime(1899, 12, 31)],
)
def test_extreme_dates_do_not_raise_raw_exception(tmp_path, value):
    out = tmp_path / "o.xlsx"
    try:
        export(out, ["Код", "v"], [["a", value]])
    except AppError:
        return
    assert _read(out)["B2"].value is not None


def test_text_longer_than_excel_cell_limit_gives_app_error_or_keeps_text(tmp_path):
    """Предел ячейки Excel — 32767 знаков: либо понятная ошибка, либо текст целиком."""
    out = tmp_path / "o.xlsx"
    big = "я" * 40000
    try:
        export(out, ["Код", "v"], [["a", big]])
    except AppError:
        return
    assert _read(out)["B2"].value == big


def test_title_longer_than_cell_limit_does_not_raise_raw(tmp_path):
    out = tmp_path / "o.xlsx"
    try:
        export(out, ["Код", "v"], [["a", 1]], title="т" * 40000)
    except AppError:
        return


def test_invalid_xml_chars_everywhere_are_removed_not_fatal(tmp_path):
    out = tmp_path / "o.xlsx"
    bad = "\ud800￾￿\x00\x0b"
    export(
        out,
        [f"Код{bad}", f"v{bad}"],
        [[f"к{bad}1", f"т{bad}ест"], [f"к2", 5]],
        title=f"Таблица{bad}",
        header_blocks=[[(("k",), f"Код{bad}")], [(("s", 0), f"v{bad}")]],
    )
    ws = _read(out)
    assert ws["A1"].value == "Таблица"
    assert ws["A3"].value == "к1" and ws["B3"].value == "тест", (ws["A3"].value, ws["B3"].value)


def test_title_made_only_of_invalid_chars_is_not_a_junk_row(tmp_path):
    """Название из одних недопустимых знаков: строка 1 пустая, а шапка не съезжает."""
    out = tmp_path / "o.xlsx"
    export(out, ["Код", "v"], [["a", 1]], title="￾\x00")
    ws = _read(out)
    heads = [(c.row, c.value) for row in ws.iter_rows() for c in row if c.value == "Код"]
    assert heads and ws["A1"].value in (None, ""), (ws["A1"].value, heads)


@pytest.mark.parametrize("n", [2 ** 53, 2 ** 53 + 1, -(2 ** 53) - 1, 10 ** 30, 2 ** 63])
def test_big_integers_round_trip_digit_by_digit(tmp_path, n):
    out = tmp_path / "o.xlsx"
    export(out, ["Код", "v"], [["a", n]])
    v = _read(out)["B2"].value
    assert str(v) == str(n), (v, n)


def test_big_integer_in_key_column_keeps_all_digits(tmp_path):
    out = tmp_path / "o.xlsx"
    key = "123456789012345678"
    export(out, ["Код", "v"], [[key, 1]])
    v = _read(out)["A2"].value
    assert str(v) == key, v


def test_timedelta_negative_and_huge_do_not_crash(tmp_path):
    out = tmp_path / "o.xlsx"
    export(out, ["Код", "v"], [["a", timedelta(hours=-5)], ["b", timedelta(days=100000)]])
    ws = _read(out)
    assert ws["B2"].value is not None and ws["B3"].value is not None


def test_exactly_16384_columns_are_accepted(tmp_path):
    out = tmp_path / "wide.xlsx"
    headers = ["k"] + [f"c{i}" for i in range(16383)]
    n = export(out, headers, [["a"] + [1] * 16383])
    assert n == 16384


def test_header_block_count_over_limit_is_rejected_with_blocks(tmp_path):
    out = tmp_path / "wide.xlsx"
    blocks = [[(("k",), "Код")]] + [[((i,), f"c{i}")] for i in range(16384)]
    with pytest.raises(AppError):
        export(out, ["x"] * len(blocks), [["a"] + [1] * 16384], header_blocks=blocks)


def test_row_limit_boundary_counts_title_and_depth(monkeypatch, tmp_path):
    """Предел строк считается от шапки и названия: ровно впритык — записывается,
    на одну больше — AppError (а не сырое исключение openpyxl)."""
    from excel_aggregator.core import exporter
    monkeypatch.setattr(exporter, "_MAX_ROWS", 20)
    out = tmp_path / "o.xlsx"
    # title(1) + шапка(1) + 18 строк = 20 — помещается
    export(out, ["Код", "v"], [[f"к{i}", i] for i in range(18)], title="т")
    with pytest.raises(AppError):
        export(out, ["Код", "v"], [[f"к{i}", i] for i in range(19)], title="т")
