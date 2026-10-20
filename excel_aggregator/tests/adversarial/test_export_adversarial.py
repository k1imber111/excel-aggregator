"""Красная команда, зона 3: запись итога (exporter), сценарий (app) и main.py."""

from __future__ import annotations

import os
import random
import subprocess
import sys
from datetime import date, datetime, timedelta
from itertools import combinations
from pathlib import Path

import openpyxl
import pytest

from excel_aggregator.core import load_workbook
from excel_aggregator.core.errors import AppError
from excel_aggregator.core.exporter import _regions, _resolve_header, export

PROJECT_ROOT = Path(__file__).resolve().parents[3]


# --------------------------------------------------------------------------
# вспомогательное


def _xlsx(path: Path, rows: list[list], fmt: dict[str, str] | None = None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Лист"
    for row in rows:
        ws.append(row)
    for coord, f in (fmt or {}).items():
        ws[coord].number_format = f
    wb.save(path)
    return path


def _table(n: int = 4) -> list[list]:
    return [["Код", "Значение"]] + [[f"к{i}", i * 1.5] for i in range(1, n + 1)]


def _run(script: list[str], cwd: Path, *args: str) -> subprocess.CompletedProcess:
    env = dict(os.environ, AGGREGATOR_SCRIPTED="1", PYTHONIOENCODING="utf-8")
    return subprocess.run(
        [sys.executable, str(PROJECT_ROOT / "main.py"), *args],
        input="\n".join(script) + "\n", capture_output=True, text=True,
        encoding="utf-8", cwd=cwd, env=env, timeout=120,
    )


def _clean_run(result: subprocess.CompletedProcess) -> None:
    out = result.stdout + result.stderr
    assert result.returncode == 0, out[-1500:]
    assert "Непредвиденная ошибка" not in out, out[-1500:]
    assert "Traceback" not in out, out[-1500:]


# --------------------------------------------------------------------------
# находки


def test_header_text_starting_with_equals_is_text_not_formula(tmp_path):
    """Заголовок «=Цена» уходит в ячейку шапки как формула, а не как текст."""
    out = tmp_path / "o.xlsx"
    export(out, ["Код", "=Цена", "=SUM(A1)"], [["a", 1, 2]])
    ws = openpyxl.load_workbook(out).active
    kinds = {c.coordinate: (c.value, c.data_type) for c in ws[1]}
    assert kinds["B1"] == ("=Цена", "s"), kinds
    assert kinds["C1"] == ("=SUM(A1)", "s"), kinds


def test_title_starting_with_equals_is_text_not_formula(tmp_path):
    out = tmp_path / "o.xlsx"
    export(out, ["Код", "v"], [["a", 1]], title="=1+1")
    cell = openpyxl.load_workbook(out).active["A1"]
    assert (cell.value, cell.data_type) == ("=1+1", "s")


def test_output_file_named_with_equals_makes_formula_title(tmp_path):
    """Имя итогового файла «=итог.xlsx» (допустимо в Windows) -> A1 формула."""
    src = _xlsx(tmp_path / "a.xlsx", _table())
    out = tmp_path / "=итог.xlsx"
    result = _run(["", "", "", str(out), "н", "н", "3"], tmp_path, str(src))
    _clean_run(result)
    cell = openpyxl.load_workbook(out).active["A1"]
    assert cell.data_type == "s", (cell.value, cell.data_type)


@pytest.mark.parametrize("bad", ["x￾y", "x\ud800y", "x￿y"])
def test_xml_noncharacters_do_not_leak_raw_parse_error(tmp_path, bad):
    """U+FFFE/U+FFFF/одиночный суррогат не вычищаются _clean: файл пишется битым,
    а _verify падает сырым ParseError вместо понятной ошибки или очистки."""
    out = tmp_path / "o.xlsx"
    try:
        export(out, ["Код", "v"], [["a", bad]])
    except AppError:
        return
    ws = openpyxl.load_workbook(out).active  # если записано — файл должен открываться
    assert ws["B2"].value is not None


def test_timedelta_over_24h_keeps_hours_in_format(tmp_path):
    """30 часов без формата источника показываются как 06:00:00 (HH:MM:SS без скобок)."""
    out = tmp_path / "o.xlsx"
    export(out, ["Код", "v"], [["a", timedelta(hours=30)]])
    fmt = openpyxl.load_workbook(out).active["B2"].number_format
    assert "[" in fmt, f"формат {fmt!r} обрезает 30 ч до 06:00:00"


def test_more_columns_than_excel_allows_is_rejected(tmp_path):
    """Лимит Excel — 16384 столбца; 16385 пишется молча, и Excel файл не откроет."""
    out = tmp_path / "wide.xlsx"
    headers = ["k"] + [f"c{i}" for i in range(16384)]
    with pytest.raises(AppError):
        export(out, headers, [["a"] + [1] * 16384])


def test_number_with_escaped_m_format_stays_number(tmp_path):
    """Формат «0\\ m» (метры) openpyxl считает датой: число 3 читается как 03.01.1900."""
    src = _xlsx(tmp_path / "s.xlsx", _table(), fmt={"B2": "0\\ m", "B3": "0\\ m"})
    grid = load_workbook(src).sheets[0]
    assert grid.row(1)[1] == 1.5 and not isinstance(grid.row(1)[1], datetime), grid.row(1)


# --------------------------------------------------------------------------
# выдержали (регрессионная сетка)


def test_text_prefixes_stay_text(tmp_path):
    out = tmp_path / "o.xlsx"
    vals = ["=1+1", "+1", "-1", "@x", "040", "04.61", "51.30", "  pad  ", "_x0041_", "a\nb"]
    export(out, ["Код", "v"], [[f"k{i}", v] for i, v in enumerate(vals)])
    ws = openpyxl.load_workbook(out).active
    assert [ws.cell(row=i + 2, column=2).value for i in range(len(vals))] == vals
    assert all(ws.cell(row=i + 2, column=2).data_type == "s" for i in range(len(vals)))


def test_custom_formats_roundtrip(tmp_path):
    out = tmp_path / "o.xlsx"
    fmts = ["0.0", "0%", '0.0 "руб."', "[Red]0.0;[Blue]-0.0", ";;;", "@", "General",
            "# ?/?", "0.00E+00", "#,##0.00_);[Red](#,##0.00)"]
    rows = [[f"k{i}", 1.25 if i % 2 else -3] for i in range(len(fmts))]
    export(out, ["Код", "v"], rows, formats=[[None, f] for f in fmts])
    ws = openpyxl.load_workbook(out).active
    assert [ws.cell(row=i + 2, column=2).number_format for i in range(len(fmts))] == fmts


def test_dates_bool_bigint_negative(tmp_path):
    out = tmp_path / "o.xlsx"
    rows = [["a", date(2020, 2, 29), True, 2**53, -7.5, datetime(2020, 1, 1, 13, 5)]]
    export(out, ["Код", "d", "b", "i", "n", "dt"], rows)
    ws = openpyxl.load_workbook(out).active
    assert ws["C2"].value is True and ws["D2"].value == 2**53 and ws["E2"].value == -7.5


def test_integer_above_2_pow_53_does_not_abort_build(tmp_path):
    """Целое > 2**53 (18-значный идентификатор) числом Excel хранит с потерей
    цифр, а _verify отвергал запись: сборка была невозможна.

    Решение: такое целое пишется текстом — все цифры сохраняются.
    """
    out = tmp_path / "o.xlsx"
    big = 123456789012345678
    export(out, ["Код", "v"], [["a", big]])
    assert openpyxl.load_workbook(out).active["B2"].value == str(big)


def test_zero_rows_and_single_column(tmp_path):
    out = tmp_path / "o.xlsx"
    assert export(out, ["Код"], []) == 0
    assert export(out, ["Код"], [["a"], ["b"]]) == 2
    assert export(out, ["Код", "v"], [], title="т") == 0


def test_huge_texts_respect_row_and_column_limits(tmp_path):
    out = tmp_path / "o.xlsx"
    big = "слово " * 3000
    export(out, ["Код " + big, "v"], [["a", big], ["b", "я" * 20000]],
           header_blocks=[[(("k",), "Код"), (("k",), "Код")],
                          [(("s", 0), big), (("l", 1), "x" * 20000)]])
    ws = openpyxl.load_workbook(out).active
    assert all((d.height or 0) <= 409 for d in ws.row_dimensions.values())
    assert all((d.width or 0) <= 255 for d in ws.column_dimensions.values())


def test_header_depths_and_repeated_ids_never_overlap_merges(tmp_path):
    random.seed(7)
    for it in range(400):
        ncols, depth = random.randint(1, 7), random.randint(1, 4)
        pool = [("g", i) for i in range(4)]
        blocks = []
        for j in range(ncols):
            levels = []
            for k in range(random.randint(1, depth)):
                ident = random.choice(pool + [("l", j, k)])
                levels.append((ident, random.choice(["", f"T{ident[1]}", "a b\nc"])))
            blocks.append(levels)
        regs = _regions(_resolve_header(blocks)[0])
        occupied: set = set()
        for _, k0, k1, j0, j1 in regs:
            cells = {(k, j) for k in range(k0, k1 + 1) for j in range(j0, j1 + 1)}
            assert not cells & occupied, (blocks, regs)
            occupied |= cells
    for it in range(25):
        blocks = [[(("s", 0), "S"), (("g", j % 2), "G"), (("l", j), f"c{j}")] for j in range(6)]
        random.shuffle(blocks)
        out = tmp_path / f"m{it}.xlsx"
        export(out, ["k"] * 6, [["a"] * 6], header_blocks=blocks)
        ranges = list(openpyxl.load_workbook(out).active.merged_cells.ranges)
        for a, b in combinations(ranges, 2):
            assert not (a.min_row <= b.max_row and b.min_row <= a.max_row
                        and a.min_col <= b.max_col and b.min_col <= a.max_col), (a, b)


# --------------------------------------------------------------------------
# сценарий и main.py


def test_argv_not_excel_nonexistent_and_directory(tmp_path):
    fake = tmp_path / "fake.xlsx"
    fake.write_text("это не Excel", encoding="utf-8")
    empty_dir = tmp_path / "пустая папка"
    empty_dir.mkdir()
    result = _run(["", "", "о", "д", "3"], tmp_path, str(fake), str(tmp_path / "нет.xlsx"),
                  str(empty_dir))
    _clean_run(result)
    assert "не является таблицей Excel" in result.stdout


def test_argv_path_with_spaces_quotes_and_cyrillic_then_output_in_new_dir(tmp_path):
    src = _xlsx(tmp_path / "Файл «отчёт» 'x' y.xlsx", _table())
    out = tmp_path / "новая папка" / "итог 1.xlsx"
    result = _run(["", "", "", f'"{out}"', "н", "н", "3"], tmp_path, str(src))
    _clean_run(result)
    assert openpyxl.load_workbook(out).active["B4"].value == 1.5


def test_output_over_source_is_refused_and_source_intact(tmp_path):
    src = _xlsx(tmp_path / "a b.xlsx", _table())
    before = src.read_bytes()
    out = tmp_path / "итог.xlsx"
    result = _run(["", "", "", str(src).upper(), str(src.with_suffix(".xls")), str(out),
                   "н", "н", "3"], tmp_path, str(src))
    _clean_run(result)
    assert src.read_bytes() == before and out.exists()


def test_declined_overwrite_keeps_existing_file(tmp_path):
    src = _xlsx(tmp_path / "a.xlsx", _table())
    out = tmp_path / "итог.xlsx"
    out.write_bytes(b"old")
    other = tmp_path / "другой.xlsx"
    result = _run(["", "", "", str(out), "", str(other), "н", "н", "3"], tmp_path, str(src))
    _clean_run(result)
    assert out.read_bytes() == b"old" and other.exists()


def test_invalid_output_name_gives_message_not_crash(tmp_path):
    src = _xlsx(tmp_path / "a.xlsx", _table())
    result = _run(["", "", "", str(tmp_path / "a?b*.xlsx"), "д", str(tmp_path / "ok.xlsx"),
                   "н", "н", "3"], tmp_path, str(src))
    _clean_run(result)
    assert (tmp_path / "ok.xlsx").exists()


def test_restart_builds_second_file_from_scratch(tmp_path):
    src = _xlsx(tmp_path / "a.xlsx", _table())
    src2 = _xlsx(tmp_path / "b.xlsx", _table(2))
    o1, o2 = tmp_path / "o1.xlsx", tmp_path / "o2.xlsx"
    result = _run(["", "", "", str(o1), "н", "д", str(src2), "", "", "", str(o2), "н", "н", "3"],
                  tmp_path, str(src))
    _clean_run(result)
    assert openpyxl.load_workbook(o1).active.max_row == 3 + 4
    assert openpyxl.load_workbook(o2).active.max_row == 3 + 2


def test_back_cancel_build_at_every_step_do_not_crash(tmp_path):
    src = _xlsx(tmp_path / "a.xlsx", _table())
    script = ["н", "с", "", "н", "с", "н", "н", "о", "д", "1", str(src), "", "", "о", "д", "3"]
    result = _run(script, tmp_path, str(src))
    _clean_run(result)
    assert "Файл собран" not in result.stdout
