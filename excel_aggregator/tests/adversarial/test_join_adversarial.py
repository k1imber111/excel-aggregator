"""Красная команда, зона 2: соединение таблиц по названиям объектов.

Сквозной путь как в программе: load_workbook -> AggregatorApp._analyse ->
Aggregator.build. Значения в источниках кодируют (источник, объект, столбец),
поэтому чужая ячейка в итоге видна сразу.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from excel_aggregator.app import AggregatorApp
from excel_aggregator.core import (
    Aggregator,
    Orientation,
    SheetGrid,
    SourceConfig,
    header_cells,
    load_workbook,
)
from excel_aggregator.core.keys import extract_keys
from excel_aggregator.core.matcher import suggest_join
from excel_aggregator.core.normalize import fold_lookalikes, normalize_key


def val(src: int, idx: int, col: int) -> int:
    """Значение, однозначно говорящее, чья это ячейка."""
    return src * 100000 + idx * 10 + col


def write_xlsx(path: Path, sheets: dict[str, list[list]], merges: dict[str, list[str]] | None = None) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        for rng in (merges or {}).get(name, []):
            ws.merge_cells(rng)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def simple_table(src: int, names: list[str], ids: dict[str, int], cols: int = 2,
                 title: str = "Название") -> list[list]:
    """Шапка + строки: название и cols числовых столбцов."""
    head = [title] + [f"Параметр {c}" for c in range(1, cols + 1)]
    return [head] + [[n] + [val(src, ids[n], c) for c in range(1, cols + 1)] for n in names]


def pipeline(paths: list[Path], fold: bool = False):
    """Как в test_samples_oracle: возвращает (app, tables, agg, headers, rows)."""
    app = AggregatorApp()
    app.books = [load_workbook(p) for p in paths]
    app._analyse()
    app.fold = fold
    tables = app._enabled()
    agg = Aggregator(app._key_title(), match=app._match)
    for t in tables:
        columns = app._data_columns(t)
        agg.add_source(
            t.keyset,
            SourceConfig(
                t.path, t.sheet.name, (t.layout.data_start, t.key_col),
                Orientation.VERTICAL, t.key_col, columns,
                [t.title(c) for c in columns], [t.paths[c] for c in columns],
                header_cells(t.grid, t.layout, columns),
                header_cells(t.grid, t.layout, [t.key_col])[0],
            ),
            t.grid,
        )
    headers, rows, _ = agg.build()
    return app, tables, agg, headers, rows


def objects(n: int, prefix: str = "Мачта") -> list[str]:
    return [f"{prefix} {i}" for i in range(1, n + 1)]


def check_rows(rows, expected: dict[str, list[list]]):
    """expected: ключ -> [значения источника 0, источника 1, ...]."""
    by_key = {r[0]: r for r in rows}
    for key, blocks in expected.items():
        assert key in by_key, f"объект {key!r} потерян; есть {list(by_key)[:12]}"
        flat = [v for block in blocks for v in block]
        assert by_key[key][1:] == flat, f"{key}: {by_key[key][1:]} != {flat}"


# ---------------------------------------------------------------- съезд строк

def test_shuffled_order_and_gaps_between_files(tmp_path):
    names = objects(10)
    ids = {n: i for i, n in enumerate(names, 1)}
    b = [names[i] for i in (9, 3, 0, 7, 5, 1, 8, 2)]  # другой порядок, пропуски 4 и 6
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": simple_table(1, names, ids)})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": simple_table(2, b, ids)})
    _, _, _, _, rows = pipeline([pa, pb])
    exp = {}
    for n in names:
        exp[n] = [[val(1, ids[n], 1), val(1, ids[n], 2)],
                  [val(2, ids[n], 1), val(2, ids[n], 2)] if n in b else [None, None]]
    check_rows(rows, exp)


def test_blank_rows_and_blank_key_inside_data(tmp_path):
    names = objects(8)
    ids = {n: i for i, n in enumerate(names, 1)}
    ta = simple_table(1, names, ids)
    ta.insert(3, [None, None, None])                    # пустая строка внутри данных
    ta.insert(6, [None, val(1, 99, 1), val(1, 99, 2)])  # данные без названия
    tb = simple_table(2, names[::-1], ids)
    tb.insert(2, [None, None, None])
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": tb})
    _, _, _, _, rows = pipeline([pa, pb])
    exp = {n: [[val(1, ids[n], 1), val(1, ids[n], 2)], [val(2, ids[n], 1), val(2, ids[n], 2)]]
           for n in names}
    check_rows(rows, exp)


def test_merged_key_cells_do_not_shift_neighbours(tmp_path):
    names = objects(8)
    ids = {n: i for i, n in enumerate(names, 1)}
    ta = simple_table(1, names, ids)
    # объединены названия двух строк (Мачта 3 и Мачта 4)
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta}, {"Л": ["A4:A5"]})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": simple_table(2, names, ids)})
    _, _, agg, _, rows = pipeline([pa, pb])
    by_key = {r[0]: r for r in rows}
    for n in names:
        if n in ("Мачта 3", "Мачта 4"):
            continue
        assert by_key[n][1:] == [val(1, ids[n], 1), val(1, ids[n], 2), val(2, ids[n], 1), val(2, ids[n], 2)]
    lost = by_key.get("Мачта 4", [None] * 3)[1] != val(1, 4, 1)
    assert not lost or agg.duplicate_warnings(), "данные строки из объединённого названия потеряны без предупреждения"


def test_real_multiword_object_without_data_is_not_dropped(tmp_path):
    names = objects(6)
    ids = {n: i for i, n in enumerate(names, 1)}
    ta = simple_table(1, names, ids)
    ta.insert(3, ["Опора 77", None, None])  # настоящий объект, характеристик пока нет
    tb = simple_table(2, names, ids)
    tb.insert(2, ["Опора 77", None, None])
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": tb})
    app, tables, _, _, rows = pipeline([pa, pb])
    # Решение: строка из одного названия в несколько слов без характеристик
    # неотличима от подписи «район 2» и в итог не идёт, но не теряется молча —
    # она в списке пропущенных подписей, который показывает шаг 2.
    assert all(
        "Опора 77" in [text for _, text in t.keyset.skipped] for t in tables
    )
    assert any("Опора 77" in w for w in app._warnings(tables[0], 6, 2))


def test_object_named_with_total_prefix_is_not_a_caption(tmp_path):
    names = objects(6) + ["Total Tower", "Всего Мачта"]
    ids = {n: i for i, n in enumerate(names, 1)}
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": simple_table(1, names, ids)})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": simple_table(2, names, ids)})
    _, _, _, _, rows = pipeline([pa, pb])
    keys = {r[0] for r in rows}
    assert {"Total Tower", "Всего Мачта"} <= keys, f"объекты с данными потеряны как подписи: {keys}"


# -------------------------------------------------------------------- ключи

def test_key_forms_join_across_types():
    forms = [2, 2.0, "2", " 2 ", "2 ", "\n2\n", "2,0", "2.0"]
    assert {normalize_key(v) for v in forms} == {"2"}
    assert normalize_key("Мачта 1") == normalize_key("Мачта\n1") == normalize_key("Мачта   1")
    assert normalize_key(1e3) == normalize_key("1000") == "1000"


def test_invisible_characters_do_not_split_keys():
    base = normalize_key("Мачта 1")
    assert normalize_key("Мачта 1​") == base      # нулевая ширина на конце
    assert normalize_key("﻿Мачта 1") == base      # BOM в начале
    assert normalize_key("Мачта​ 1") == base


def test_huge_digit_string_does_not_crash():
    # int() ограничен 4300 цифрами: ячейка из цифр не должна ронять программу
    assert normalize_key("1" * 5000) is not None


def test_very_long_integers_keep_identity():
    assert normalize_key(12345678901234567890) == "12345678901234567890"
    assert normalize_key("12345678901234567890") == "12345678901234567890"
    assert normalize_key(1e22) == normalize_key("10000000000000000000000")


def test_lookalikes_fold_only_on_request():
    assert normalize_key("A1") != normalize_key("А1")  # без согласия не склеиваем
    assert fold_lookalikes("A1") == fold_lookalikes("А1")
    assert fold_lookalikes("Мачта X") == fold_lookalikes("Мaчтa Х")


def test_fold_does_not_silently_drop_distinct_object_inside_source():
    # В одном файле два РАЗНЫХ объекта «A1» (лат.) и «А1» (кир.) с разными данными.
    grid = SheetGrid("s", [["Имя", "v"], ["A1", 1], ["А1", 2], ["B2", 3]])
    ks = extract_keys(grid, 1, 0, Orientation.VERTICAL)
    assert ks.keys == ["A1", "А1", "B2"]
    agg = Aggregator("Имя", match=fold_lookalikes)
    cfg = SourceConfig(Path("f.xlsx"), "s", (1, 0), Orientation.VERTICAL, 0, [1], ["v"])
    agg.add_source(ks, cfg, grid)
    _, rows = agg.build_table()
    values = sorted(r[1] for r in rows if r[1] is not None)
    assert values == [1, 2, 3] or agg.duplicate_warnings(), (
        f"значение 2 пропало без предупреждения: {rows}"
    )


def test_numeric_cell_key_matches_text_key_across_files(tmp_path):
    ta = [["Номер", "П1", "П2"]] + [[i, val(1, i, 1), val(1, i, 2)] for i in range(1, 9)]
    tb = [["Номер", "П1", "П2"]] + [[(f" {i} " if i % 2 else f"{i}.0"), val(2, i, 1), val(2, i, 2)]
                                   for i in range(8, 0, -1)]
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": tb})
    _, _, _, _, rows = pipeline([pa, pb])
    assert len(rows) == 8, [r[0] for r in rows]
    for r in rows:
        i = int(r[0])
        assert r[1:] == [val(1, i, 1), val(1, i, 2), val(2, i, 1), val(2, i, 2)]


def test_case_and_whitespace_variants_join(tmp_path):
    names = objects(8)
    ids = {n: i for i, n in enumerate(names, 1)}
    ta = simple_table(1, names, ids)
    tb = simple_table(2, names, ids)
    tb[1][0] = "МАЧТА 1"
    tb[2][0] = "мачта\n2"
    tb[3][0] = "  Мачта   3 "
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": tb})
    _, _, _, _, rows = pipeline([pa, pb])
    assert len(rows) == 8, [r[0] for r in rows]
    for r in rows:
        i = ids[r[0]]
        assert r[1:] == [val(1, i, 1), val(1, i, 2), val(2, i, 1), val(2, i, 2)]


def test_caption_between_data_rows_is_skipped_and_neighbours_not_shifted(tmp_path):
    names = objects(8)
    ids = {n: i for i, n in enumerate(names, 1)}
    ta = simple_table(1, names, ids)
    ta.insert(4, ["Район 2", None, None])
    ta.insert(7, ["Итого 5", 10, 20])
    ta.insert(9, ["Примечание: см. приложение", None, None])
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": simple_table(2, names, ids)})
    _, _, _, _, rows = pipeline([pa, pb])
    keys = [r[0] for r in rows]
    assert keys == names, keys
    check_rows(rows, {n: [[val(1, ids[n], 1), val(1, ids[n], 2)],
                          [val(2, ids[n], 1), val(2, ids[n], 2)]] for n in names})


# ----------------------------------------------------- выбор столбца соединения

def test_names_not_in_first_column(tmp_path):
    names = objects(8)
    ids = {n: i for i, n in enumerate(names, 1)}
    ta = [["Код", "Название", "П"]] + [["K%d" % (i + 50), n, val(1, ids[n], 1)] for i, n in enumerate(names)]
    tb = [["Название", "Тип", "П"]] + [[n, "T", val(2, ids[n], 1)] for n in names[::-1]]
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": tb})
    _, tables, _, _, rows = pipeline([pa, pb])
    assert [t.key_col for t in tables] == [1, 0]
    assert len(rows) == 8
    for r in rows:
        assert val(1, ids[r[0]], 1) in r and val(2, ids[r[0]], 1) in r


def test_row_number_columns_do_not_join_partially_overlapping_tables(tmp_path):
    # Два каталога по 20 объектов с нумерацией 1..20; общих названий всего 5
    common = objects(5)
    a_names = common + [f"Опора {i}" for i in range(1, 16)]
    b_names = common + [f"Антенна {i}" for i in range(1, 16)]
    ids = {n: i for i, n in enumerate(dict.fromkeys(a_names + b_names), 1)}

    def table(src, ns):
        return [["№ п/п", "Название", "Высота", "Масса"]] + [
            [i, n, val(src, ids[n], 1), val(src, ids[n], 2)] for i, n in enumerate(ns, 1)
        ]
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": table(1, a_names)})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": table(2, b_names)})
    _, tables, _, _, rows = pipeline([pa, pb])
    assert [t.key_col for t in tables] == [1, 1], f"соединили по столбцам {[t.key_col for t in tables]}"
    assert len(rows) == 35, len(rows)


def test_tables_without_common_names_are_not_mixed_by_numbers(tmp_path):
    na = [f"Опора {i}" for i in range(1, 9)]
    nb = [f"Антенна {i}" for i in range(1, 9)]
    ids = {n: i for i, n in enumerate(na + nb, 1)}
    ta = [["№", "Название", "П"]] + [[i, n, val(1, ids[n], 1)] for i, n in enumerate(na, 1)]
    tb = [["№", "Название", "П"]] + [[i, n, val(2, ids[n], 1)] for i, n in enumerate(nb, 1)]
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": tb})
    _, tables, _, _, rows = pipeline([pa, pb])
    assert [t.key_col for t in tables] == [1, 1]
    assert len(rows) == 16


def test_integer_measure_column_does_not_beat_names(tmp_path):
    # Высоты (целые, все различны) в обоих файлах те же, названия пересекаются частично
    heights = [120, 135, 150, 165, 180, 195, 210, 225]
    a_names = ["Мачта 1", "Мачта 2", "Мачта 3", "Опора 1", "Опора 2", "Опора 3", "Опора 4", "Опора 5"]
    b_names = ["Мачта 1", "Мачта 2", "Мачта 3", "Башня 1", "Башня 2", "Башня 3", "Башня 4", "Башня 5"]
    ids = {n: i for i, n in enumerate(dict.fromkeys(a_names + b_names), 1)}
    ta = [["Название", "Высота", "П"]] + [[n, h, val(1, ids[n], 1)] for n, h in zip(a_names, heights)]
    tb = [["Название", "Высота", "П"]] + [[n, h, val(2, ids[n], 1)] for n, h in zip(b_names, heights)]
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": tb})
    _, tables, _, _, rows = pipeline([pa, pb])
    assert [t.key_col for t in tables] == [0, 0], [t.key_col for t in tables]


def test_unrelated_third_source_does_not_poison_join(tmp_path):
    names = objects(8)
    other = [f"Река {i}" for i in range(1, 9)]
    ids = {n: i for i, n in enumerate(names, 1)}
    ids.update({n: i + 50 for i, n in enumerate(other, 1)})
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": simple_table(1, names, ids)})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": simple_table(2, names[::-1], ids)})
    pc = write_xlsx(tmp_path / "c.xlsx", {"Л": simple_table(3, other, ids)})
    _, _, _, _, rows = pipeline([pa, pb, pc])
    assert len(rows) == 16
    by_key = {r[0]: r for r in rows}
    for n in names:
        assert by_key[n][1:] == [val(1, ids[n], 1), val(1, ids[n], 2), val(2, ids[n], 1),
                                 val(2, ids[n], 2), None, None]
    for n in other:
        assert by_key[n][1:] == [None] * 4 + [val(3, ids[n], 1), val(3, ids[n], 2)]


def test_chain_join_three_sources_and_union_order(tmp_path):
    n1 = objects(6)
    n2 = objects(6)[3:] + [f"Опора {i}" for i in range(1, 4)]
    n3 = [f"Опора {i}" for i in range(1, 4)] + [f"Антенна {i}" for i in range(1, 4)]
    ids = {n: i for i, n in enumerate(dict.fromkeys(n1 + n2 + n3), 1)}
    ps = [write_xlsx(tmp_path / f"{k}.xlsx", {"Л": simple_table(k, ns, ids)})
          for k, ns in ((1, n1), (2, n2), (3, n3))]
    _, _, agg, _, rows = pipeline(ps)
    assert [r[0] for r in rows] == list(dict.fromkeys(n1 + n2 + n3))
    for r in rows:
        i = ids[r[0]]
        for k, ns in ((1, n1), (2, n2), (3, n3)):
            got = r[1 + 2 * (k - 1): 3 + 2 * (k - 1)]
            assert got == ([val(k, i, 1), val(k, i, 2)] if r[0] in ns else [None, None])
    cov = {name.split(" / ")[0]: c for name, c, _ in agg.coverage_stats()}
    assert cov == {"1.xlsx": 6, "2.xlsx": 6, "3.xlsx": 6}


def test_two_sheets_of_one_file_both_join(tmp_path):
    names = objects(8)
    ids = {n: i for i, n in enumerate(names, 1)}
    pa = write_xlsx(tmp_path / "a.xlsx", {
        "Север": simple_table(1, names[:5], ids),
        "Юг": simple_table(2, names[3:], ids),
    })
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": simple_table(3, names[::-1], ids)})
    _, tables, agg, _, rows = pipeline([pa, pb])
    assert len(tables) == 3
    labels = agg.source_labels()
    assert len(set(labels)) == 3, labels
    by_key = {r[0]: r for r in rows}
    assert len(rows) == 8
    for n in names:
        i = ids[n]
        assert by_key[n][5:7] == [val(3, i, 1), val(3, i, 2)]
        north = [val(1, i, 1), val(1, i, 2)] if n in names[:5] else [None, None]
        south = [val(2, i, 1), val(2, i, 2)] if n in names[3:] else [None, None]
        assert by_key[n][1:3] == north and by_key[n][3:5] == south


# ----------------------------------------------- повторы, подписи, header_blocks

def test_duplicate_in_source_first_occurrence_with_warning(tmp_path):
    names = objects(6)
    ids = {n: i for i, n in enumerate(names, 1)}
    ta = simple_table(1, names, ids)
    ta.append(["мачта 2", 7777, 8888])
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": simple_table(2, names, ids)})
    _, _, agg, _, rows = pipeline([pa, pb])
    by_key = {r[0]: r for r in rows}
    assert by_key["Мачта 2"][1:3] == [val(1, 2, 1), val(1, 2, 2)]
    w = agg.duplicate_warnings()
    assert len(w) == 1 and "мачта 2" in w[0] and "строка 8" in w[0]


def test_same_file_name_from_different_folders_keeps_blocks_apart(tmp_path):
    names = objects(6)
    ids = {n: i for i, n in enumerate(names, 1)}
    p1 = write_xlsx(tmp_path / "x" / "каталог.xlsx", {"Л": simple_table(1, names, ids)})
    p2 = write_xlsx(tmp_path / "y" / "каталог.xlsx", {"Л": simple_table(2, names, ids)})
    _, _, agg, _, _ = pipeline([p1, p2])
    assert len(set(agg.source_labels())) == 2
    blocks = agg.header_blocks()
    assert len({b[0][0] for b in blocks[1:]}) == 2   # id верхнего уровня различны
    assert len({b[0][1] for b in blocks[1:]}) == 2   # подписи различны


def test_labels_collide_with_numbered_stem(tmp_path):
    names = objects(6)
    ids = {n: i for i, n in enumerate(names, 1)}
    ps = [
        write_xlsx(tmp_path / "x" / "каталог.xlsx", {"Л": simple_table(1, names, ids)}),
        write_xlsx(tmp_path / "y" / "каталог.xlsx", {"Л": simple_table(2, names, ids)}),
        write_xlsx(tmp_path / "z" / "каталог (2).xlsx", {"Л": simple_table(3, names, ids)}),
    ]
    _, _, agg, _, _ = pipeline(ps)
    labels = agg.source_labels()
    assert len(set(labels)) == 3, labels


def test_source_without_selected_columns_in_header_blocks():
    g1 = SheetGrid("s", [["Имя", "a", "b"], ["x y", 1, 2], ["z w", 3, 4]])
    g2 = SheetGrid("t", [["Имя", "c"], ["x y", 5], ["q r", 6]])
    agg = Aggregator("Имя")
    for g, cols in ((g1, [1, 2]), (g2, [])):
        ks = extract_keys(g, 1, 0, Orientation.VERTICAL)
        agg.add_source(ks, SourceConfig(Path(f"{g.name}.xlsx"), g.name, (1, 0), Orientation.VERTICAL,
                                        0, cols, [str(c) for c in cols]), g)
    headers, rows = agg.build_table()
    assert len(headers) == 3
    assert [r[0] for r in rows] == ["x y", "z w", "q r"]
    blocks = agg.header_blocks()
    assert blocks is None or len(blocks) == len(headers)


def test_suggest_join_single_table_and_empty_inputs():
    assert suggest_join([]) == []
    g = SheetGrid("s", [["Имя", "v"], ["a b", 1], ["c d", 2], ["e f", 3]])
    assert suggest_join([(g, 1)]) == [0]
    empty = SheetGrid("e", [[None, None], [None, None]])
    assert suggest_join([(g, 1), (empty, 1)])[1] is None


# ------------------------------------------------------------- вторая пачка

def test_unrelated_first_source_does_not_hide_join_of_others(tmp_path):
    other = [f"Река {i}" for i in range(1, 9)]
    names = objects(8)
    ids = {n: i for i, n in enumerate(names, 1)}
    ids.update({n: i + 50 for i, n in enumerate(other, 1)})
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": simple_table(1, other, ids)})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": simple_table(2, names, ids)})
    pc = write_xlsx(tmp_path / "c.xlsx", {"Л": simple_table(3, names[::-1], ids)})
    _, _, _, _, rows = pipeline([pa, pb, pc])
    assert len(rows) == 16
    by_key = {r[0]: r for r in rows}
    for n in names:
        assert by_key[n][3:] == [val(2, ids[n], 1), val(2, ids[n], 2), val(3, ids[n], 1), val(3, ids[n], 2)]


def test_real_objects_named_like_district_with_data_are_kept(tmp_path):
    names = [f"район {i}" for i in range(1, 9)]
    ids = {n: i for i, n in enumerate(names, 1)}
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": simple_table(1, names, ids)})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": simple_table(2, names[::-1], ids)})
    _, _, _, _, rows = pipeline([pa, pb])
    assert [r[0] for r in rows] == names
    for r in rows:
        i = ids[r[0]]
        assert r[1:] == [val(1, i, 1), val(1, i, 2), val(2, i, 1), val(2, i, 2)]


def test_fold_lookalikes_joins_latin_and_cyrillic_through_pipeline(tmp_path):
    cyr = ["Мачта А1", "Мачта В2", "Мачта С3", "Мачта Е4", "Мачта К5", "Мачта М6", "Мачта Н7", "Мачта О8"]
    lat = ["Мачта A1", "Мачта B2", "Мачта C3", "Мачта E4", "Мачта K5", "Мачта M6", "Мачта H7", "Мачта O8"]
    ids = {n: i for i, n in enumerate(cyr, 1)}
    ids.update({n: i for i, n in enumerate(lat, 1)})
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": simple_table(1, cyr, ids)})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": simple_table(2, lat, ids)})
    _, _, _, _, rows_no = pipeline([pa, pb], fold=False)
    assert len(rows_no) == 16            # без согласия не склеиваем
    _, _, _, _, rows = pipeline([pa, pb], fold=True)
    assert len(rows) == 8
    for r in rows:
        i = ids[r[0]]
        assert r[1:] == [val(1, i, 1), val(1, i, 2), val(2, i, 1), val(2, i, 2)]


def test_numeric_keys_with_text_forms_and_long_ids_through_pipeline(tmp_path):
    big = [10**12 + i for i in range(1, 9)]            # длинные целые (идентификаторы)
    ta = [["ID", "П1"]] + [[v, val(1, i, 1)] for i, v in enumerate(big, 1)]
    tb = [["ID", "П1"]] + [[str(v), val(2, i, 1)] for i, v in reversed(list(enumerate(big, 1)))]
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": tb})
    _, _, _, _, rows = pipeline([pa, pb])
    assert len(rows) == 8, [r[0] for r in rows]
    for r in rows:
        i = big.index(int(r[0])) + 1
        assert r[1:] == [val(1, i, 1), val(2, i, 1)]


def test_table_with_names_in_late_column_and_small_overlap_still_joins(tmp_path):
    names = objects(4)
    ids = {n: i for i, n in enumerate(names, 1)}
    # названия в 5-м столбце, пересечение только 2 объекта из 4
    ta = [["a", "b", "c", "d", "Название"]] + [
        [val(1, i, 1) + k for k in range(4)] + [n] for i, n in enumerate(names[:3], 1)
    ]
    tb = [["Название", "x"]] + [[n, val(2, ids[n], 1)] for n in names[1:]]
    pa = write_xlsx(tmp_path / "a.xlsx", {"Л": ta})
    pb = write_xlsx(tmp_path / "b.xlsx", {"Л": tb})
    _, tables, _, _, rows = pipeline([pa, pb])
    assert [t.key_col for t in tables] == [4, 0]
    assert len(rows) == 4
