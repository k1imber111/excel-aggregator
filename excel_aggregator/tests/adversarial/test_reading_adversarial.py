"""Красная команда, зона 1: чтение листа и распознавание таблицы."""

from __future__ import annotations

import datetime as dt

import openpyxl
import pytest

from excel_aggregator.core.layout import (
    column_paths,
    detect_layout,
    header_cells,
    header_text,
    layout_from,
)
from excel_aggregator.core.models import SheetGrid
from excel_aggregator.core.reader import load_workbook

N = None


def G(rows, merged=None, formats=None):
    return SheetGrid("t", rows, merged or [], formats)


# --- layout: распознавание ---------------------------------------------------

@pytest.mark.parametrize("ncols", [3, 4])
def test_vertical_merge_in_data_does_not_hide_rows(ncols):
    """Объединённый по вертикали «Район» в узкой таблице не должен выкидывать строки данных."""
    if ncols == 3:
        rows = [
            ["Район", "Объект", "Значение"],
            ["Север", "Объект 1", 10], [N, "Объект 2", 20], [N, "Объект 3", 30],
            ["Юг", "Объект 4", 40], [N, "Объект 5", 50],
        ]
        merged = [(1, 4, 0, 1), (4, 6, 0, 1)]
    else:
        rows = [
            ["Район", "Объект", "Значение", "Группа"],
            ["Север", "Объект 1", 10, "A"], [N, "Объект 2", 20, N], [N, "Объект 3", 30, N],
            ["Юг", "Объект 4", 40, "B"], [N, "Объект 5", 50, N],
        ]
        merged = [(1, 4, 0, 1), (4, 6, 0, 1), (1, 4, 3, 4), (4, 6, 3, 4)]
    layout = detect_layout(G(rows, merged))
    assert layout is not None
    assert layout.data_start == 1, f"data_start={layout.data_start}, header_rows={layout.header_rows}"


def test_rowspan_header_of_single_column_is_not_title():
    """«Объект» объединён по 2 строкам, остальные заголовки во 2-й строке: это шапка, а не подпись."""
    rows = [
        ["Объект", N, N],
        [N, "X", "Y"],
        ["О1", 1, 2], ["О2", 3, 4], ["О3", 5, 6],
    ]
    grid = G(rows, [(0, 2, 0, 1)])
    layout = detect_layout(grid)
    assert layout.title == "", f"title={layout.title!r}, header_rows={layout.header_rows}"
    assert column_paths(grid, layout)[0] == ["Объект"]


def test_group_over_half_of_columns_is_not_swallowed_as_title():
    """Группа «Координаты» над 4 из 6 столбцов (A и F без верхнего уровня) — часть шапки."""
    rows = [
        [N, "Координаты", N, N, N, N],
        ["Объект", "X", "Y", "Z", "W", "Прим"],
        ["О1", 1, 2, 3, 4, "п"], ["О2", 5, 6, 7, 8, "т"], ["О3", 9, 1, 2, 3, "у"],
    ]
    grid = G(rows, [(0, 1, 1, 5)])
    paths = column_paths(grid, detect_layout(grid))
    assert paths[1] == ["Координаты", "X"], paths


def test_first_data_row_with_only_text_note_is_data():
    """Первый объект заполнен частично (имя + текстовое примечание) — не шапка."""
    rows = [
        ["Наименование", "X", "Y", "Z", "Прим."],
        ["Объект А", N, N, N, "примечание"],
        ["Б", 1, 2, 3, "т"], ["В", 4, 5, 6, "т"], ["Г", 7, 8, 9, "т"],
    ]
    layout = detect_layout(G(rows))
    assert layout.header_rows == [0], f"header_rows={layout.header_rows}, data_start={layout.data_start}"


def test_small_ints_first_data_row_is_not_numbering():
    """Строка данных 1, 2, 3 (счёт, не нумерация граф) под шапкой не теряется."""
    rows = [
        ["Ключ", "A", "B"],
        [1, 2, 3], [4, 5, 6], [7, 8, 9],
    ]
    layout = detect_layout(G(rows))
    assert layout.data_start == 1, f"data_start={layout.data_start}, numbering={layout.numbering_row}"


def test_title_merged_over_table_with_far_stray_cell_stays_title():
    """Подпись, объединённая на ширину таблицы, при одиночной ячейке далеко справа остаётся подписью."""
    stray = [N] * 40
    stray[0] = "Г"
    stray[39] = "заметка"
    rows = [
        ["Каталог 2024"] + [N] * 39,
        ["Объект", "X", "Y", "Z", "W", "V"] + [N] * 34,
        ["О1", 1, 2, 3, 4, 5] + [N] * 34,
        ["О2", 1, 2, 3, 4, 5] + [N] * 34,
        stray,
    ]
    grid = G(rows, [(0, 1, 0, 6)])
    layout = detect_layout(grid)
    assert layout.title == "Каталог 2024", f"title={layout.title!r}, header_rows={layout.header_rows}"


def test_unknown_unit_in_unit_row_is_not_dropped():
    """«кг» в строке единиц («м», «фут», «кг») — не подпись раздела; заголовок не пропадает."""
    rows = [
        ["Объект", "Высота", N, "Вес"],
        [N, "м", "фут", "кг"],
        ["О1", 1, 2, 3], ["О2", 4, 5, 6], ["О3", 7, 8, 9],
    ]
    grid = G(rows)
    layout = detect_layout(grid)
    paths = column_paths(grid, layout)
    assert any("кг" in p for p in paths[3]), paths
    assert "кг" in " ".join(header_text(grid, layout, r, 3) for r in layout.header_rows)


def test_table_not_from_a1_with_notes_around():
    rows = [
        ["Справка", N, N, N, N],
        [N, N, N, N, N],
        [N, "Объект", "X", "Y", N],
        [N, "О1", 1, 2, N], [N, "О2", 3, 4, N], [N, "О3", 5, 6, N],
        ["Составил: Иванов", N, N, N, N],
    ]
    grid = G(rows)
    layout = detect_layout(grid)
    assert layout.data_start == 3 and layout.header_rows == [2]
    assert column_paths(grid, layout)[2] == ["X"]


# --- layout: устойчивость и тексты -------------------------------------------

@pytest.mark.parametrize("rows", [
    [[1]], [["a", "b"]], [["a"], ["b"]], [[N, N], [N, N]], [],
    [[1, 2], [3]], [["a", N, N], [N, N, "b"]],
])
def test_degenerate_grids_do_not_raise(rows):
    grid = G(rows)
    layout = detect_layout(grid)
    if layout is not None:
        column_paths(grid, layout)
        header_cells(grid, layout, list(range(grid.ncols)))


def test_merge_wider_than_sheet_does_not_raise():
    rows = [
        ["Объект", "Группа", N],
        [N, "X", "Y"],
        ["О1", 1, 2], ["О2", 3, 4], ["О3", 5, 6],
    ]
    grid = G(rows, [(0, 1, 1, 10), (0, 2, 0, 1)])
    layout = detect_layout(grid)
    paths = column_paths(grid, layout)
    assert paths[1][0] == "Группа" and paths[2][0] == "Группа"
    cells = header_cells(grid, layout, [0, 1, 2])
    assert len(cells) == 3


def test_header_whitespace_and_newlines_normalised():
    grid = G([["  Высота   над\n\n уровнем  ", "X"], ["a", 1], ["b", 2], ["c", 3]])
    layout = layout_from(grid, 1)
    assert header_text(grid, layout, 0, 0) == "Высота над\nуровнем"
    assert column_paths(grid, layout)[0] == ["Высота над уровнем"]


def test_header_cells_nonadjacent_and_headerless_column():
    rows = [
        [N, "Группа", N, N, "Прим", N],
        ["Объект", "X", "Y", "Z", N, N],
        ["О1", 1, 2, 3, 9, 7], ["О2", 4, 5, 6, 9, 7], ["О3", 7, 8, 9, 9, 7],
    ]
    grid = G(rows, [(0, 1, 1, 4)])
    layout = layout_from(grid, 2)
    cells = header_cells(grid, layout, [3, 0, 5])
    assert cells[0] == [((0, 1), "Группа"), ((1, 3), "Z")]
    assert cells[1][-1] == ((1, 0), "Объект")
    assert cells[2][-1][1] == "Столбец F"


def test_vertical_merge_in_headerless_numeric_table():
    rows = [
        ["Север", 1, 10], [N, 2, 20], [N, 3, 30],
        ["Юг", 4, 40], [N, 5, 50],
    ]
    grid = G(rows, [(0, 3, 0, 1), (3, 5, 0, 1)])
    layout = detect_layout(grid)
    assert layout is not None and layout.data_start == 0, layout


# --- SheetGrid.transposed ----------------------------------------------------

def test_transposed_keeps_formats_merges_and_flags():
    grid = G(
        [["a", "b", "c"], [1, 2, 3]],
        [(0, 1, 0, 2)],
        [[N, N, N], ["0.0", "0%", "dd.mm.yyyy"]],
    )
    grid.hidden = True
    grid.error_cells = 2
    t = grid.transposed()
    assert (t.nrows, t.ncols) == (3, 2)
    assert t.number_format(2, 1) == "dd.mm.yyyy" and t.number_format(1, 1) == "0%"
    assert t.span(0, 0) == (0, 2, 0, 1)
    assert t.value(1, 0) == "a" and t.raw(1, 0) is None
    assert t.hidden and t.error_cells == 2


def test_transposed_formats_with_ragged_rows():
    grid = G([["a", "b", "c"], [1, 2, 3]], None, [["x"], ["0.0", "0.0", "0.0"]])
    t = grid.transposed()
    assert t.number_format(2, 1) == "0.0" and t.number_format(0, 0) == "x"


# --- reader: xlsx ------------------------------------------------------------

def _save(tmp_path, build, name="a.xlsx"):
    wb = openpyxl.Workbook()
    build(wb)
    path = tmp_path / name
    wb.save(path)
    return path


def test_reader_formats_and_types(tmp_path):
    def build(wb):
        ws = wb.active
        ws.append(["Объект", "A", "B", "C", "D", "E"])
        ws.append(["040", 1.5, 0.25, dt.datetime(2024, 5, 1), 3, "51.30"])
        ws["B2"].number_format = "0.0"
        ws["C2"].number_format = "0%"
        ws["D2"].number_format = "dd.mm.yyyy"
    grid = load_workbook(_save(tmp_path, build)).sheets[0]
    assert grid.value(1, 0) == "040" and grid.value(1, 5) == "51.30"
    assert isinstance(grid.value(1, 3), dt.datetime)
    assert [grid.number_format(1, c) for c in (1, 2, 3, 4)] == ["0.0", "0%", "dd.mm.yyyy", None]


def test_reader_custom_format_preserved(tmp_path):
    def build(wb):
        ws = wb.active
        ws.append(["Объект", "A"])
        ws.append(["x", 3])
        ws["B2"].number_format = '0.00" км"'
    grid = load_workbook(_save(tmp_path, build)).sheets[0]
    assert grid.number_format(1, 1) == '0.00" км"'


def test_reader_errors_spaces_formula(tmp_path):
    def build(wb):
        ws = wb.active
        ws.append(["Объект", "A", "B"])
        ws.append(["x", "#REF!", "   "])
        ws.append(["y", "=1+1", "\xa0\t"])
    grid = load_workbook(_save(tmp_path, build)).sheets[0]
    assert grid.error_cells == 1
    assert grid.uncached_formulas == 1
    assert grid.value(1, 1) is None and grid.value(1, 2) is None
    assert grid.value(2, 1) is None and grid.ncols == 3


def test_reader_zero_width_cell_does_not_widen_table(tmp_path):
    """Ячейка из одного невидимого символа (U+200B) — пустая, не должна растягивать таблицу."""
    def build(wb):
        ws = wb.active
        ws.append(["Объект", "A"])
        ws.append(["x", 1])
        ws["Z5"] = "​"
    grid = load_workbook(_save(tmp_path, build)).sheets[0]
    assert (grid.nrows, grid.ncols) == (2, 2)


def test_reader_hidden_empty_single_cell_sheets(tmp_path):
    def build(wb):
        ws = wb.active
        ws.title = "vis"
        ws["A1"] = 5
        h = wb.create_sheet("hid")
        h["A1"] = "a"
        h.sheet_state = "hidden"
        v = wb.create_sheet("vh")
        v["A1"] = "b"
        v.sheet_state = "veryHidden"
        wb.create_sheet("empty")
    book = load_workbook(_save(tmp_path, build))
    by = {s.name: s for s in book.sheets}
    assert not by["vis"].hidden and by["hid"].hidden and by["vh"].hidden
    assert (by["vis"].nrows, by["vis"].ncols) == (1, 1)
    assert (by["empty"].nrows, by["empty"].ncols) == (0, 0)
    assert detect_layout(by["vis"]) is None and detect_layout(by["empty"]) is None


def test_reader_merge_wider_than_data_is_clipped_and_formats_follow_anchor(tmp_path):
    def build(wb):
        ws = wb.active
        ws["A1"] = "Заголовок"
        ws.merge_cells("A1:J1")
        ws.append(["Объект", "A", "B"])
        ws.append(["x", 1, 2])
        ws["A1"].number_format = "@"
    grid = load_workbook(_save(tmp_path, build)).sheets[0]
    assert grid.ncols == 3
    assert grid.span(0, 2) == (0, 1, 0, 3)
    assert grid.number_format(0, 2) == "@"


def test_reader_merged_table_pipeline_header(tmp_path):
    def build(wb):
        ws = wb.active
        ws.append(["Объект", "Координаты", None, "Высота"])
        ws.append([None, "X", "Y", "м"])
        ws.merge_cells("A1:A2")
        ws.merge_cells("B1:C1")
        for i in range(3):
            ws.append([f"О{i}", i, i + 1, i + 2])
    grid = load_workbook(_save(tmp_path, build)).sheets[0]
    layout = detect_layout(grid)
    assert layout.data_start == 2 and layout.header_rows == [0, 1]
    assert column_paths(grid, layout) == {
        0: ["Объект"], 1: ["Координаты", "X"], 2: ["Координаты", "Y"], 3: ["Высота, м"],
    }
