"""Регрессионные и сценарные тесты: приложение в памяти с заданными ответами."""

from __future__ import annotations

import datetime
import io
from pathlib import Path

import openpyxl
import pytest
from rich.console import Console

import excel_aggregator.app as app_module
from excel_aggregator.app import AggregatorApp
from excel_aggregator.core.errors import OutputWriteError
from excel_aggregator.core.exporter import export
from excel_aggregator.core.models import SheetGrid
from excel_aggregator.core.paths import resolve_input_path
from excel_aggregator.core.reader import load_workbook
from excel_aggregator.ui.widgets import UI

SAMPLE_WITH_DATES = (
    Path(__file__).resolve().parents[2] / "Примеры таблиц" / "Каталог Сочи 2024 МОС ФАП-262.xls"
)


def _make_xlsx(path: Path, rows: list[list]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


@pytest.fixture
def table(tmp_path: Path) -> Path:
    """Таблица с шапкой: названия объектов A, B, C и две графы характеристик."""
    return _make_xlsx(tmp_path / "таблица.xlsx", [
        ["Название", "Высота", "Ширина"],
        ["A", 1, 10],
        ["B", 2, 20],
        ["C", 3, 30],
    ])


def _drive(answers: list[str], monkeypatch) -> tuple[AggregatorApp, str]:
    """Прогоняет приложение в памяти с заранее заданными ответами."""
    monkeypatch.setenv("AGGREGATOR_SCRIPTED", "1")
    app = AggregatorApp()
    buf = io.StringIO()
    app.ui.console = Console(file=buf, width=200)
    queue = list(answers)
    app.ui.console.input = lambda prompt="": queue.pop(0)
    app.run()
    assert not queue, f"Остались непрочитанные ответы: {queue}"
    return app, buf.getvalue()


def _leaf_headers(path: Path) -> list:
    """Нижняя строка шапки итогового файла (та, на которой стоит фильтр)."""
    ws = openpyxl.load_workbook(path).active
    leaf = int(ws.auto_filter.ref.split(":")[0][1:])
    return [ws.cell(row=leaf, column=c).value for c in range(1, ws.max_column + 1)]


# --- сценарий -----------------------------------------------------------------

def test_shortest_path_single_file(table, tmp_path, monkeypatch):
    # файл -> Enter (к проверке) -> Enter (верно) -> Enter (все столбцы) -> имя
    out = tmp_path / "итог"
    app, text = _drive(
        ["1", str(table), "", "", "", str(out), "н", "н", "3"], monkeypatch
    )
    assert "Запись проверена" in text
    ws = openpyxl.load_workbook(tmp_path / "итог.xlsx").active
    assert ws["A1"].value == "итог"
    # строка 2 — источник; заголовок первого столбца растянут на обе строки шапки
    assert [c.value for c in ws[2]] == ["Название", "таблица", None]
    assert [c.value for c in ws[3]] == [None, "Высота", "Ширина"]
    assert "A2:A3" in {str(r) for r in ws.merged_cells.ranges}
    assert [c.value for c in ws[4]] == ["A", 1, 10]


def test_back_and_forth_does_not_duplicate_sources(table, monkeypatch):
    # проверка -> «назад» к файлам -> снова проверка: источник по-прежнему один
    app, _ = _drive(["1", str(table), "", "н", "", "о", "д", "3"], monkeypatch)
    assert len(app.tables) == 1
    assert app.tables[0].keyset.keys == ["A", "B", "C"]


def test_same_file_is_not_added_twice(table, monkeypatch):
    app, text = _drive(["1", str(table), str(table), "о", "д", "3"], monkeypatch)
    assert len(app.books) == 1
    assert "Файл уже добавлен" in text


def test_two_files_join_by_shared_names(tmp_path, monkeypatch):
    first = _make_xlsx(tmp_path / "первый.xlsx", [
        ["Код", "Высота"], ["k1", 1], ["k2", 2], ["k3", 3],
    ])
    # столбец названий — второй, а не первый; объект k9 есть только здесь
    second = _make_xlsx(tmp_path / "второй.xlsx", [
        ["Вес", "Код"], [5, "k2"], [6, "k3"], [7, "k9"],
    ])
    out = tmp_path / "итог"
    _drive(["1", str(first), str(second), "", "", "", "", str(out), "н", "н", "3"], monkeypatch)
    ws = openpyxl.load_workbook(tmp_path / "итог.xlsx").active
    assert [c.value for c in ws[2]] == ["Код", "первый", "второй"]      # строка источников
    rows = {r[0]: r[1:] for r in ws.iter_rows(min_row=4, values_only=True)}
    assert rows == {"k1": (1, None), "k2": (2, 5), "k3": (3, 6), "k9": (None, 7)}


def test_lookalike_letters_are_asked_about(tmp_path, monkeypatch):
    latin = _make_xlsx(tmp_path / "латиница.xlsx", [["Код", "Высота"], ["A1", 1], ["A2", 2]])
    cyrillic = _make_xlsx(tmp_path / "кириллица.xlsx", [["Код", "Вес"], ["А1", 5], ["А2", 6]])
    out = tmp_path / "итог"
    answers = ["1", str(latin), str(cyrillic), "", "д", "", "", "", str(out), "н", "н", "3"]
    _, text = _drive(answers, monkeypatch)
    assert "разными буквами" in text
    ws = openpyxl.load_workbook(tmp_path / "итог.xlsx").active
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    assert rows == [("A1", 1, 5), ("A2", 2, 6)]      # два объекта, а не четыре


def test_table_without_header_gets_generic_titles(tmp_path, monkeypatch):
    bare = _make_xlsx(tmp_path / "без шапки.xlsx", [["A", 1], ["B", 2], ["C", 3]])
    out = tmp_path / "итог"
    _drive(["1", str(bare), "", "", "", str(out), "н", "н", "3"], monkeypatch)
    ws = openpyxl.load_workbook(tmp_path / "итог.xlsx").active
    assert ws["A2"].value == "Столбец A"          # заголовок первого столбца — с верхней строки
    assert _leaf_headers(tmp_path / "итог.xlsx")[1] == "Столбец B"


def test_manual_data_start_and_key_column(tmp_path, monkeypatch):
    # Пользователь поправляет найденное: данные со строки 3, названия — столбец B
    odd = _make_xlsx(tmp_path / "таблица.xlsx", [
        ["Высота", "Название"], [1, "A"], [2, "B"], [3, "C"],
    ])
    answers = [
        "1", str(odd), "",
        "2", "1",            # изменить источник № 1
        "1", "3",            # первая строка данных -> 3
        "2", "2",            # столбец названий -> второй в списке (B)
        "",                  # готово
        "о", "д", "3",
    ]
    app, _ = _drive(answers, monkeypatch)
    source = app.tables[0]
    assert source.layout.data_start == 2
    assert source.key_col == 1
    assert source.keyset.keys == ["B", "C"]


def test_tables_without_shared_names_are_reported_not_joined(tmp_path, monkeypatch):
    first = _make_xlsx(tmp_path / "первый.xlsx", [["Код", "Вес"], ["a1", 1], ["a2", 2], ["a3", 3]])
    second = _make_xlsx(tmp_path / "второй.xlsx", [["Код", "Рост"], ["b1", 1], ["b2", 2], ["b3", 3]])
    app, text = _drive(["1", str(first), str(second), "", "о", "д", "3"], monkeypatch)
    assert [t.key_col for t in app.tables] == [0, 0]     # не «Вес» и «Рост»
    assert "Общих названий объектов между таблицами не найдено" in text


def test_source_file_cannot_be_overwritten_by_result(table, tmp_path, monkeypatch):
    # Имя итога совпало с исходным файлом — программа просит другое имя
    safe = tmp_path / "итог"
    answers = ["1", str(table), "", "", "", str(table), str(safe), "н", "н", "3"]
    _, text = _drive(answers, monkeypatch)
    assert "один из исходных файлов" in text
    assert (tmp_path / "итог.xlsx").is_file()
    # исходник цел: в нём по-прежнему простая шапка в первой строке
    assert openpyxl.load_workbook(table).active["A1"].value == "Название"


def test_output_name_with_dot_is_not_truncated(table, tmp_path, monkeypatch):
    target = tmp_path / "отчёт v1.2"
    _drive(["1", str(table), "", "", "", str(target), "н", "н", "3"], monkeypatch)
    assert (tmp_path / "отчёт v1.2.xlsx").is_file()


# --- страховка: непредвиденная ошибка не выбрасывает из приложения ----------

def test_unexpected_error_is_logged_and_offers_retry(tmp_path, monkeypatch):
    monkeypatch.setattr(app_module, "_program_dir", lambda: tmp_path)

    def boom(self):
        raise RuntimeError("сбой шага")

    monkeypatch.setattr(AggregatorApp, "_step_files", boom)
    _, out = _drive(["1", "н", "3"], monkeypatch)   # начать, не повторять, выход
    assert "Непредвиденная ошибка" in out
    assert "сбой шага" in (tmp_path / "aggregator_error.log").read_text(encoding="utf-8")


# --- чтение и экспорт ----------------------------------------------------------

@pytest.mark.skipif(not SAMPLE_WITH_DATES.is_file(), reason="Пример недоступен")
def test_xls_date_cells_are_dates():
    grid = load_workbook(SAMPLE_WITH_DATES).sheet("Т12")
    assert grid.value(0, 1) == datetime.datetime(2018, 4, 24)


@pytest.mark.skipif(not SAMPLE_WITH_DATES.is_file(), reason="Пример недоступен")
def test_error_cells_are_counted():
    book = load_workbook(SAMPLE_WITH_DATES)
    assert sum(sheet.error_cells for sheet in book.sheets) == 946


def test_uncached_formulas_are_counted(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Название", "Сумма"])
    ws.append(["A", "=1+1"])        # файл сохранён не из Excel: значения формулы нет
    wb.save(tmp_path / "формулы.xlsx")
    grid = load_workbook(tmp_path / "формулы.xlsx").sheets[0]
    assert grid.uncached_formulas == 1
    assert grid.value(1, 1) is None


def test_export_keeps_dates_and_strips_illegal_characters(tmp_path):
    out = tmp_path / "итог.xlsx"
    when = datetime.datetime(2024, 1, 15)
    export(out, ["Ключ", "Да\x0bта"], [["k\x01", when]], key_title="Ключ")
    ws = openpyxl.load_workbook(out).active
    assert ws["B1"].value == "Дата"
    assert ws["A2"].value == "k"
    assert ws["B2"].value == when


def test_export_unwritable_folder_is_output_write_error(tmp_path):
    blocker = tmp_path / "это_файл"
    blocker.write_text("x", encoding="utf-8")
    with pytest.raises(OutputWriteError):
        export(blocker / "папка" / "итог.xlsx", ["A"], [["1"]], key_title="A")


def test_input_name_with_dot_gets_extension(tmp_path):
    f = tmp_path / "данные.2024.xls"
    f.write_bytes(b"\xd0\xcf\x11\xe0")
    assert resolve_input_path(str(tmp_path / "данные.2024")) == f


# --- rich-разметка в данных ----------------------------------------------------

def test_brackets_in_data_are_printed_literally():
    ui = UI()
    buf = io.StringIO()
    ui.console = Console(file=buf, width=200)
    ui.info("Таблица [final].xls и хвост [/]")
    ui.error("значение [/b]")
    ui.show_preview(SheetGrid("лист [a]", [["x [/b]", "y [red]"]]), title="Лист [a]")
    ui.show_table("Сводка [x]", ["Файл"], [["имя [final].xls"]])
    out = buf.getvalue()
    for fragment in ("[final]", "хвост [/]", "значение [/b]", "x [/b]", "y [red]",
                     "Лист [a]", "Сводка [x]", "имя [final].xls"):
        assert fragment in out
