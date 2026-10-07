"""Тесты экспортёра: форматирование, типы, ведущие нули."""

import openpyxl
import pytest

from excel_aggregator.core.errors import OutputWriteError
from excel_aggregator.core import exporter
from excel_aggregator.core.exporter import export


def _read_back(path):
    wb = openpyxl.load_workbook(path)
    return wb.active


def test_export_basic(tmp_path):
    headers = ["№", "Название", "Высота"]
    rows = [["2", "Трубы котельной", 1412.18], ["3", "Мачта", 1517.93]]
    out = tmp_path / "вложенная" / "итог.xlsx"   # родительская папка создаётся
    export(out, headers, rows, key_title="№")

    ws = _read_back(out)
    assert [c.value for c in ws[1]] == headers
    assert ws.max_row == 3
    assert ws.max_column == 3
    assert ws.freeze_panes == "B2"   # шапка и первый столбец
    assert ws.auto_filter.ref == "A1:C3"
    # шапка оформлена
    assert ws["A1"].font.bold
    assert ws["A1"].fill.fgColor.rgb in ("001F4E78", "FF1F4E78")
    assert ws["B2"].alignment.wrap_text
    # числа остались числами
    assert isinstance(ws["C2"].value, float)


def test_export_leading_zeros_stay_text(tmp_path):
    headers = ["Код", "Долгота", "Число"]
    rows = [["040", "04.61", 4634]]
    out = tmp_path / "итог.xlsx"
    export(out, headers, rows, key_title="Код")

    ws = _read_back(out)
    assert ws["A2"].value == "040"          # нули не потеряны
    assert ws["A2"].number_format == "@"
    assert ws["B2"].value == "04.61"
    assert ws["B2"].number_format == "@"
    assert ws["C2"].value == 4634           # настоящее число — числом
    assert isinstance(ws["C2"].value, int)


def test_export_column_widths_bounded(tmp_path):
    headers = ["A", "B"]
    rows = [["x", "y" * 200]]
    out = tmp_path / "итог.xlsx"
    export(out, headers, rows, key_title="A")
    ws = _read_back(out)
    assert ws.column_dimensions["A"].width >= 10
    assert ws.column_dimensions["B"].width <= 50


def test_export_row_height_grows_for_wrapped_text(tmp_path):
    headers = ["A", "B"]
    rows = [["x", "строка 1\nстрока 2\nстрока 3"]]
    out = tmp_path / "итог.xlsx"
    export(out, headers, rows, key_title="A")
    ws = _read_back(out)
    assert ws["B2"].alignment.wrap_text
    assert ws.row_dimensions[2].height > 15


def test_export_permission_error(tmp_path, monkeypatch):
    out = tmp_path / "итог.xlsx"

    def deny(*args, **kwargs):
        raise PermissionError("занят")

    monkeypatch.setattr(openpyxl.Workbook, "save", deny)
    with pytest.raises(OutputWriteError) as excinfo:
        export(out, ["A"], [["1"]], key_title="A")
    assert "закройте" in excinfo.value.message


def test_export_multilevel_header(tmp_path):
    paths = [
        ["№"],
        ["Файл 1", "Координаты", "X"],
        ["Файл 1", "Координаты", "Y"],
        ["Файл 2", "Высота, м"],
    ]
    rows = [["2", 1.5, 2.5, 10], ["040", 3.5, 4.5, None]]
    out = tmp_path / "итог.xlsx"
    checked = export(out, [], rows, key_title="№", header_paths=paths, title="Сводная")

    ws = _read_back(out)
    assert checked == 8                              # все ячейки данных сверены с файлом
    assert ws["A1"].value == "Сводная"               # название таблицы
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert {"A2:A4", "B2:C2", "B3:C3", "D3:D4"} <= merged
    assert ws["B2"].value == "Файл 1" and ws["D2"].value == "Файл 2"   # строка источников
    assert ws["B4"].value == "X" and ws["C4"].value == "Y"
    assert ws["D3"].value == "Высота, м"             # короткая цепочка растянута вниз
    assert ws.auto_filter.ref == "A4:D6"             # фильтр на нижней строке шапки
    assert ws.freeze_panes == "B5"
    assert ws["B2"].fill.fgColor.rgb != ws["D2"].fill.fgColor.rgb     # блоки источников различимы
    assert ws["A5"].value == 2 and isinstance(ws["A5"].value, int)    # название-число — числом
    assert ws["A6"].value == "040"                   # ведущий ноль — текстом


def test_export_text_starting_with_equals_is_not_a_formula(tmp_path):
    out = tmp_path / "итог.xlsx"
    export(out, ["A", "B"], [["k", "=1+1"]], key_title="A")
    cell = _read_back(out)["B2"]
    assert cell.value == "=1+1" and cell.data_type == "s"


def test_verify_detects_difference(tmp_path):
    out = tmp_path / "итог.xlsx"
    export(out, ["A", "B"], [["k", 1]], key_title="A")
    with pytest.raises(OutputWriteError):
        exporter._verify(out, [["k", 2]], data_row=2, ncols=2)


def test_verify_tolerates_float_precision_of_xlsx(tmp_path):
    # 17-я значащая цифра не помещается в .xlsx — это не расхождение
    out = tmp_path / "итог.xlsx"
    assert export(out, ["A", "B"], [["k", 41189.053493908126]], key_title="A") == 2
