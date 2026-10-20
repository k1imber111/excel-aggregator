"""Запись итоговой таблицы в форматированный .xlsx и проверка записанного."""

from __future__ import annotations

import math
import re
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook as _XlsxWorkbook
from openpyxl import load_workbook as _load_xlsx
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from .errors import OutputWriteError

# Число с ведущим нулём в целой части: "040", "04.61", "-05" — сохраняем как текст
_LEADING_ZERO_RE = re.compile(r"-?0\d")
# Название объекта из одних цифр («22») — пишем числом, чтобы фильтр сортировал 1, 2, 10
_INT_KEY_RE = re.compile(r"-?(0|[1-9]\d{0,14})")

_KEY_FILL = PatternFill("solid", fgColor="1F4E78")
# Источники чередуют заливку, чтобы блоки столбцов разных файлов различались
_SOURCE_FILLS = (PatternFill("solid", fgColor="2E75B6"), _KEY_FILL)
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14, color="1F4E78")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
# Подпись широкого блока — слева: по центру блока в 15 столбцов она
# оказалась бы за краем экрана
_WIDE_ALIGNMENT = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
_WIDE_GROUP = 6       # группа шире стольких столбцов подписывается слева
_GENERAL_DIGITS = 11  # столько знаков показывает Excel у числа в общем формате
_HEADER_LINE = 12     # строка заголовка длиннее — переносится, а не расширяет столбец
_BOLD_FACTOR = 1.1    # жирный текст шапки шире обычного
_CELL_ALIGNMENT = Alignment(vertical="center", wrap_text=True)
_THIN_SIDE = Side(style="thin")
_THIN_BORDER = Border(left=_THIN_SIDE, right=_THIN_SIDE, top=_THIN_SIDE, bottom=_THIN_SIDE)

_MIN_WIDTH = 8
_MAX_WIDTH = 50
_DEFAULT_ROW_HEIGHT = 15
_MAX_ROW_HEIGHT = 409
_TITLE_ROW_HEIGHT = 30

_DATE_FORMAT = "DD.MM.YYYY"
_DATETIME_FORMAT = "DD.MM.YYYY HH:MM"
_TIME_FORMAT = "HH:MM:SS"
_DURATION_FORMAT = "[HH]:MM:SS"   # длительность от суток и больше не обрезается

# Пределы листа Excel
_MAX_COLUMNS = 16384
_MAX_ROWS = 1048576
# Точнее Excel целое не хранит (число двойной точности)
_MAX_EXACT_INT = 2 ** 53
# Суррогаты и U+FFFE/U+FFFF недопустимы в XML, а ILLEGAL_CHARACTERS_RE их пропускает
_XML_INVALID_RE = re.compile(
    "[%s-%s%s%s]" % (chr(0xD800), chr(0xDFFF), chr(0xFFFE), chr(0xFFFF))
)

_AS_IS = (bool, int, float, datetime, date, time, timedelta)


def _clean(text) -> str:
    """Текст, допустимый в .xlsx: без управляющих символов, переводы строк — «\\n»."""
    text = _XML_INVALID_RE.sub("", ILLEGAL_CHARACTERS_RE.sub("", str(text)))
    return text.replace("\r\n", "\n").replace("\r", "\n")


def _cell_value(value, is_key: bool = False):
    """Значение, которое попадёт в ячейку: числа и даты — как есть, остальное — текст.

    Целое длиннее 2**53 (18-значный код) Excel числом не сохранит без потери
    цифр — оно пишется текстом.
    """
    if isinstance(value, int) and not isinstance(value, bool) and abs(value) > _MAX_EXACT_INT:
        return str(value)
    if value is None or isinstance(value, _AS_IS):
        return value
    text = _clean(value)
    if is_key and _INT_KEY_RE.fullmatch(text):
        return int(text)
    return text


def _cell_format(value, fmt: str | None) -> str | None:
    """Формат числа для записи: из исходника; дате без формата — русский вид."""
    if isinstance(value, bool) or not isinstance(value, _AS_IS):
        return None
    if fmt:
        return fmt
    if isinstance(value, datetime):
        return _DATETIME_FORMAT if value.time() != time() else _DATE_FORMAT
    if isinstance(value, date):
        return _DATE_FORMAT
    if isinstance(value, timedelta):
        return _DURATION_FORMAT
    if isinstance(value, time):
        return _TIME_FORMAT
    return None


def _number_text(value: float, fmt: str | None) -> str:
    """Примерный вид числа в Excel при формате fmt — для подбора ширины."""
    if fmt is None:
        if float(value).is_integer() and abs(value) < 1e15:
            return str(int(value))
        return f"{value:.{_GENERAL_DIGITS - 1}g}"[:_GENERAL_DIGITS + 1]
    section = fmt.split(";")[0]
    literals = sum(len(m) for m in re.findall(r'"([^"]*)"', section))
    section = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.|[_*].', "", section)
    if not re.search(r"[0#?]", section):
        return section or str(value)
    if "%" in section:
        value *= 100
    if "E" in section.upper():
        return f"{value:.2E}"
    decimals = re.search(r"\.([0#?]+)", section)
    places = len(decimals.group(1)) if decimals else 0
    integer = section.split(".")[0]
    text = f"{abs(value):,.{places}f}" if "," in integer else f"{abs(value):.{places}f}"
    whole, dot, fraction = text.partition(".")
    whole = whole.zfill(integer.count("0"))   # формат «000»: 40 -> 040
    text = ("-" if value < 0 else "") + whole + dot + fraction
    return text + ("%" if "%" in section else "") + " " * literals


def _display_text(value, fmt: str | None) -> str:
    """Как значение примерно выглядит в ячейке (для ширины и высоты)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "ИСТИНА" if value else "ЛОЖЬ"
    if isinstance(value, (int, float)):
        return _number_text(value, fmt)
    if isinstance(value, (datetime, date, time, timedelta)):
        shown = fmt or _cell_format(value, None) or ""
        return re.sub(r'"([^"]*)"|\[[^\]]*\]|\\', r"\1", shown.split(";")[0]) or str(value)
    return str(value)


def _wrap_lines(text: str, width: float, factor: float = 1.0) -> int:
    """Сколько строк займёт текст в ячейке ширины width при переносе по словам."""
    usable = max(1, int((width - 1) / factor))
    total = 0
    for line in (text or "").split("\n"):
        lines, length = 1, 0
        for word in line.split(" "):
            if length and length + 1 + len(word) > usable:
                lines, length = lines + 1, 0
            length += len(word) + (1 if length else 0)
            while length > usable:          # слово длиннее строки рвётся
                lines, length = lines + 1, length - usable
        total += lines
    return total


def _same(expected, actual) -> bool:
    """Совпадает ли прочитанное из файла значение с записанным."""
    if expected is None or expected == "":
        return actual is None or actual == ""
    if isinstance(expected, datetime):
        return isinstance(actual, datetime) and abs((actual - expected).total_seconds()) < 0.001
    if isinstance(expected, date):
        return (actual.date() if isinstance(actual, datetime) else actual) == expected
    if isinstance(expected, (time, timedelta)):
        return actual is not None  # Excel хранит время долей суток — сверяем наличие
    if isinstance(expected, float) and isinstance(actual, (int, float)):
        # .xlsx хранит 15–16 значащих цифр: 17-я цифра исходного числа теряется
        return math.isclose(actual, expected, rel_tol=1e-12, abs_tol=0.0)
    return actual == expected


def _verify(
    path: Path,
    rows: list[list],
    data_row: int,
    ncols: int,
    formats: list[list[str | None]] | None = None,
) -> int:
    """Перечитывает записанный файл и сверяет данные и форматы чисел ячейка в ячейку.

    Возвращает число сверенных ячеек; расхождение -> OutputWriteError.
    """
    try:
        book = _load_xlsx(path, read_only=True)
    except Exception as exc:
        raise OutputWriteError(
            f"Файл «{path}» записан, но не открывается для проверки: {exc}"
        ) from exc
    try:
        written = [
            [(cell.value, getattr(cell, "number_format", None)) for cell in row]
            for row in book.active.iter_rows(
                min_row=data_row, max_row=data_row + len(rows) - 1, max_col=ncols
            )
        ] if rows else []
    except Exception as exc:  # строки читаются лениво: битый XML всплывает здесь
        raise OutputWriteError(
            f"Файл «{path}» записан, но не читается для проверки: {exc}"
        ) from exc
    finally:
        book.close()
    if len(written) != len(rows):
        raise OutputWriteError(
            f"Проверка записи не пройдена: в файле {len(written)} строк данных, "
            f"ожидалось {len(rows)}."
        )
    checked = 0
    for r, (expected_row, actual_row) in enumerate(zip(rows, written)):
        for c, value in enumerate(expected_row):
            expected = _cell_value(value, c == 0)
            actual, actual_format = actual_row[c] if c < len(actual_row) else (None, None)
            where = f"{get_column_letter(c + 1)}{data_row + r}"
            if not _same(expected, actual):
                raise OutputWriteError(
                    f"Проверка записи не пройдена: ячейка {where} — в файле {actual!r}, "
                    f"ожидалось {expected!r}."
                )
            fmt = formats[r][c] if formats and c < len(formats[r]) else None
            expected_format = _cell_format(expected, fmt)
            if expected_format and actual_format != expected_format:
                raise OutputWriteError(
                    f"Проверка записи не пройдена: формат ячейки {where} — в файле "
                    f"«{actual_format}», ожидался «{expected_format}»."
                )
            checked += 1
    return checked


def _blocks_from_paths(paths: list[list[str]]) -> list[list[tuple[object, str]]]:
    """Шапка-блоки из цепочек названий: группа — соседние столбцы с общим началом цепочки."""
    return [
        [
            (("group", tuple(path[:k + 1])) if k < len(path) - 1 else ("leaf", j), text)
            for k, text in enumerate(path)
        ]
        for j, path in enumerate(paths)
    ]


def _resolve_header(blocks: list[list[tuple[object, str]]]) -> tuple[list[list], dict]:
    """Матрица id ячеек шапки [уровень][столбец] и тексты по id.

    Короткая шапка источника растягивается вниз; пустая ячейка присоединяется
    к заголовку того же столбца сверху или снизу, если тот занимает один
    столбец. Нулевой уровень — подписи источников — не трогается.
    """
    depth = max(len(levels) for levels in blocks)
    texts: dict = {}
    ids = [[None] * len(blocks) for _ in range(depth)]
    for j, levels in enumerate(blocks):
        for k in range(depth):
            ident, text = levels[min(k, len(levels) - 1)]
            ids[k][j] = ident
            texts.setdefault(ident, _clean(text))
    ncols = len(blocks)

    def empty(k: int, j: int) -> bool:
        return not texts[ids[k][j]]

    def one_column(k: int, j: int) -> bool:
        ident = ids[k][j]
        return not (j > 0 and ids[k][j - 1] == ident) and not (
            j + 1 < ncols and ids[k][j + 1] == ident
        )

    for j in range(ncols):
        for k in range(2, depth):              # пусто под заголовком — к нему
            if empty(k, j) and not any(not empty(b, j) for b in range(k + 1, depth)) and (
                not empty(k - 1, j) and one_column(k - 1, j)
            ):
                ids[k][j] = ids[k - 1][j]
        for k in range(depth - 2, 0, -1):      # пусто над заголовком — к нему
            if empty(k, j) and not any(not empty(a, j) for a in range(1, k)) and (
                not empty(k + 1, j) and one_column(k + 1, j)
            ):
                ids[k][j] = ids[k + 1][j]
    return ids, texts


def _regions(ids: list[list]) -> list[tuple[object, int, int, int, int]]:
    """Области шапки (id, k0, k1, j0, j1) — включительно; не прямоугольник — по ячейкам."""
    cells: dict = {}
    for k, row in enumerate(ids):
        for j, ident in enumerate(row):
            cells.setdefault(ident, []).append((k, j))
    result = []
    for ident, members in cells.items():
        ks = [k for k, _ in members]
        js = [j for _, j in members]
        k0, k1, j0, j1 = min(ks), max(ks), min(js), max(js)
        if (k1 - k0 + 1) * (j1 - j0 + 1) == len(members):
            result.append((ident, k0, k1, j0, j1))
        else:
            result.extend((ident, k, k, j, j) for k, j in members)
    return result


def export(
    path: Path,
    headers: list[str],
    rows: list[list],
    key_title: str = "Название",
    header_paths: list[list[str]] | None = None,
    title: str | None = None,
    header_blocks: list[list[tuple[object, str]]] | None = None,
    formats: list[list[str | None]] | None = None,
) -> int:
    """Пишет итоговую таблицу в .xlsx с оформлением и проверяет записанное.

    header_blocks — шапка как в исходниках: на каждый столбец уровни
    (id, текст), одинаковый id у соседних ячеек — одна объединённая ячейка
    (см. Aggregator.header_blocks). header_paths — упрощённая шапка цепочками
    [источник, группа…, название]. Без них шапка — одна строка из headers.
    title — строка названия таблицы над шапкой. formats — форматы чисел
    ячеек rows («0.0», «DD.MM.YYYY»; None — общий).

    - шапка: заливка (блоки разных источников чередуют цвет), белый жирный
      текст, переносы строк как в исходнике, центрирование;
    - данные: тонкие границы, вертикальное выравнивание по центру, форматы
      чисел и дат из исходников;
    - ширина столбцов по показу данных и заголовков (мин 8, макс 50), высота
      строк — по переносу текста;
    - закреплены шапка и первый столбец, автофильтр на нижней строке шапки;
    - строки с ведущими нулями — текстом ("@");
    - управляющие символы, недопустимые в .xlsx, из текста удаляются;
    - печать: альбомная, по ширине страницы, шапка на каждом листе.

    Возвращает число ячеек данных, сверенных с файлом после записи.
    """
    path = Path(path)
    banded = header_blocks is not None or header_paths is not None
    if header_blocks is None:
        paths = header_paths or [[h] for h in (headers or [key_title])]
        header_blocks = _blocks_from_paths([[str(level) for level in p] or [""] for p in paths])
    ncols_total = len(header_blocks)
    if ncols_total > _MAX_COLUMNS:
        raise OutputWriteError(
            f"В итоге {ncols_total} столбцов, а лист Excel вмещает не больше {_MAX_COLUMNS}. "
            "Выберите меньше столбцов на шаге 3."
        )
    if (1 if title else 0) + len(max(header_blocks, key=len)) + len(rows) > _MAX_ROWS:
        raise OutputWriteError(
            f"В итоге {len(rows)} строк, а лист Excel вмещает не больше {_MAX_ROWS}."
        )
    ids, texts = _resolve_header(header_blocks)
    regions = _regions(ids)
    ncols = len(header_blocks)
    depth = len(ids)
    top = 2 if title else 1            # первая строка шапки
    leaf_row = top + depth - 1         # нижняя строка шапки — на ней фильтр
    data_row = leaf_row + 1
    last_col = get_column_letter(ncols)

    wb = _XlsxWorkbook()
    ws = wb.active
    ws.title = "Итог"

    if title:
        # Без объединения: текст сам продолжается вправо по пустым ячейкам строки
        cell = ws.cell(row=1, column=1, value=_clean(title))
        cell.data_type = "s"  # название «=итог» — текст, не формула
        cell.font = _TITLE_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = _TITLE_ROW_HEIGHT

    # --- данные ---------------------------------------------------------
    values = [[_cell_value(v, c == 0) for c, v in enumerate(row)] for row in rows]
    shown: list[list[str]] = []
    for r, row in enumerate(values):
        row_formats = formats[r] if formats and r < len(formats) else []
        shown_row = []
        for c, value in enumerate(row):
            cell = ws.cell(row=data_row + r, column=c + 1, value=value)
            fmt = _cell_format(value, row_formats[c] if c < len(row_formats) else None)
            if isinstance(value, str):
                cell.data_type = "s"  # текст, начинающийся с «=», — не формула
                if _LEADING_ZERO_RE.match(value):
                    cell.number_format = "@"  # не терять ведущие нули
            elif fmt:
                cell.number_format = fmt
            cell.border = _THIN_BORDER
            cell.alignment = _CELL_ALIGNMENT
            shown_row.append(_display_text(value, fmt))
        shown.append(shown_row)

    # --- ширина столбцов ----------------------------------------------------
    widths: list[float] = []
    for j in range(ncols):
        data_len = max(
            (max(len(line) for line in row[j].split("\n")) for row in shown if j < len(row)),
            default=0,
        )
        head_len = 0
        for ident, k0, k1, j0, j1 in regions:
            # подпись источника (только нулевой уровень) ширину не задаёт
            if j0 == j1 == j and k1 > 0 and texts[ident]:
                lines = texts[ident].split("\n")
                longest_word = max(len(w) for line in lines for w in line.split(" ") or [""])
                longest_line = max(len(line) for line in lines)
                head_len = max(head_len, longest_word * _BOLD_FACTOR,
                               min(longest_line, _HEADER_LINE) * _BOLD_FACTOR)
        widths.append(min(_MAX_WIDTH, max(_MIN_WIDTH, max(data_len, head_len) + 2)))
    for ident, k0, k1, j0, j1 in regions:
        if j1 > j0 and texts[ident]:
            # Группа над несколькими столбцами: самое длинное слово должно поместиться
            need = max(len(w) for w in texts[ident].replace("\n", " ").split(" ")) * _BOLD_FACTOR + 2
            span = sum(widths[j0:j1 + 1])
            if span < need:
                widths[j1] = min(_MAX_WIDTH, widths[j1] + need - span)
    for j, width in enumerate(widths):
        ws.column_dimensions[get_column_letter(j + 1)].width = round(width, 1)

    for r, row in enumerate(shown):
        lines = max(
            (_wrap_lines(text, widths[c]) for c, text in enumerate(row)
             if isinstance(values[r][c], str)),
            default=1,
        )
        if lines > 1:
            ws.row_dimensions[data_row + r].height = min(
                _MAX_ROW_HEIGHT, lines * _DEFAULT_ROW_HEIGHT
            )

    # --- шапка ------------------------------------------------------------
    fills = []
    band, previous = -1, None
    for j in range(ncols):
        if banded and j > 0:
            if ids[0][j] != previous:
                band, previous = band + 1, ids[0][j]
            fills.append(_SOURCE_FILLS[band % len(_SOURCE_FILLS)])
        else:
            fills.append(_KEY_FILL)
    for k in range(depth):
        for j in range(ncols):
            cell = ws.cell(row=top + k, column=j + 1)
            cell.fill = fills[j]
            cell.font = _HEADER_FONT
            cell.alignment = _HEADER_ALIGNMENT
            cell.border = _THIN_BORDER

    heights = [_DEFAULT_ROW_HEIGHT] * depth
    tall: list[tuple[int, int, float]] = []   # (первый уровень, последний, нужная высота)
    for ident, k0, k1, j0, j1 in regions:
        text = texts[ident]
        if not text:
            continue
        cell = ws.cell(row=top + k0, column=j0 + 1, value=text)
        cell.data_type = "s"  # заголовок «=Цена» — текст, не формула
        if j1 - j0 + 1 > _WIDE_GROUP:
            cell.alignment = _WIDE_ALIGNMENT
        if k1 > k0 or j1 > j0:
            ws.merge_cells(
                start_row=top + k0, start_column=j0 + 1, end_row=top + k1, end_column=j1 + 1
            )
        need = _wrap_lines(text, sum(widths[j0:j1 + 1]), _BOLD_FACTOR) * _DEFAULT_ROW_HEIGHT
        if k1 > k0:
            tall.append((k0, k1, need))
        else:
            heights[k0] = max(heights[k0], need)
    for k0, k1, need in sorted(tall, key=lambda t: t[1] - t[0]):
        heights[k1] += max(0, need - sum(heights[k0:k1 + 1]))
    for k, height in enumerate(heights):
        ws.row_dimensions[top + k].height = min(_MAX_ROW_HEIGHT, height)

    ws.freeze_panes = ws.cell(row=data_row, column=2 if ncols > 1 else 1)
    ws.auto_filter.ref = f"A{leaf_row}:{last_col}{max(leaf_row, data_row + len(rows) - 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = f"1:{leaf_row}"

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
    except PermissionError as exc:
        raise OutputWriteError(
            f"Не удалось записать файл «{path}»: он открыт в Excel (закройте его "
            "и повторите) или в эту папку нет прав на запись."
        ) from exc
    except OSError as exc:
        raise OutputWriteError(f"Не удалось записать файл «{path}»: {exc}") from exc

    return _verify(path, rows, data_row, ncols, formats)
