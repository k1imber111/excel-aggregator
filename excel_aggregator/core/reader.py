"""Чтение книг Excel (.xls через xlrd, .xlsx через openpyxl) в SheetGrid.

Формат определяется по сигнатуре файла, а не по расширению.
Merged cells разворачиваются, ячейки-ошибки (#REF! и т.п.) -> None,
реальная область данных вычисляется по факту заполнения (DIMENSIONS не доверяем).
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import xlrd
from xlrd.biffh import XLRDError

from .errors import (
    CorruptedFileError,
    EncryptedFileError,
    NotExcelFileError,
    PathNotFoundError,
)
from .models import SheetGrid, Workbook

_OLE2_SIGNATURE = b"\xd0\xcf\x11\xe0"
_ZIP_SIGNATURE = b"PK"

# Строковые значения ячеек-ошибок Excel
_ERROR_VALUES = {
    "#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A",
    "#GETTING_DATA", "#SPILL!", "#CALC!",
}


def _is_empty(v) -> bool:
    """Пустая ячейка: None или строка из одних пробелов."""
    return v is None or (isinstance(v, str) and v.strip() == "")


def _real_area(data: list[list]) -> tuple[int, int]:
    """(nrows, ncols) реальной области по фактически заполненным ячейкам."""
    max_r = -1
    max_c = -1
    for r, row in enumerate(data):
        for c, v in enumerate(row):
            if not _is_empty(v):
                if r > max_r:
                    max_r = r
                if c > max_c:
                    max_c = c
    return max_r + 1, max_c + 1


def _build_grid(name: str, data: list[list], merged_ranges) -> SheetGrid:
    """Обрезает данные до реальной области и собирает SheetGrid."""
    nrows, ncols = _real_area(data)
    trimmed = [row[:ncols] for row in data[:nrows]]
    clipped = [
        (rlo, min(rhi, nrows), clo, min(chi, ncols))
        for rlo, rhi, clo, chi in merged_ranges
        if rlo < nrows and clo < ncols
    ]
    return SheetGrid(name, trimmed, clipped)


def _map_open_error(exc: Exception, path: Path):
    """Превращает исключения движков чтения в пользовательские ошибки."""
    text = str(exc).lower()
    if "password" in text or "encrypt" in text:
        return EncryptedFileError(f"Файл защищён паролем: {path.name}")
    return CorruptedFileError(f"Не удалось прочитать файл «{path.name}»: {exc}")


def _load_xls(path: Path) -> Workbook:
    try:
        book = xlrd.open_workbook(str(path), formatting_info=True)
    except XLRDError as exc:
        # formatting_info может мешать на экзотических файлах — пробуем без него
        try:
            book = xlrd.open_workbook(str(path))
        except XLRDError as exc2:
            raise _map_open_error(exc2, path) from exc2
        else:
            return _xls_grids(path, book)
    except Exception as exc:
        raise _map_open_error(exc, path) from exc
    return _xls_grids(path, book)


def _xls_grids(path: Path, book) -> Workbook:
    sheets = []
    for sh in book.sheets():
        data = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype in (
                    xlrd.XL_CELL_EMPTY,
                    xlrd.XL_CELL_BLANK,
                    xlrd.XL_CELL_ERROR,
                ):
                    row.append(None)
                else:
                    row.append(cell.value)
            data.append(row)
        merged = getattr(sh, "merged_cells", []) or []
        sheets.append(_build_grid(sh.name, data, merged))
    return Workbook(path=path, sheets=sheets)


def _load_xlsx(path: Path) -> Workbook:
    import openpyxl

    try:
        book = openpyxl.load_workbook(path, data_only=True, read_only=False)
    except (zipfile.BadZipFile, OSError) as exc:
        raise _map_open_error(exc, path) from exc
    except Exception as exc:
        raise _map_open_error(exc, path) from exc

    sheets = []
    for ws in book.worksheets:
        data = []
        for row in ws.iter_rows(values_only=True):
            data.append([
                None if (v is None or (isinstance(v, str) and v.strip() == "") or v in _ERROR_VALUES)
                else v
                for v in row
            ])
        merged = [
            (rng.min_row - 1, rng.max_row, rng.min_col - 1, rng.max_col)
            for rng in ws.merged_cells.ranges
        ]
        sheets.append(_build_grid(ws.title, data, merged))
    return Workbook(path=path, sheets=sheets)


def load_workbook(path: Path) -> Workbook:
    """Читает книгу Excel в Workbook. Формат — по сигнатуре файла."""
    path = Path(path)
    if not path.is_file():
        raise PathNotFoundError(f"Файл не найден: {path}")
    try:
        with open(path, "rb") as fh:
            signature = fh.read(8)
    except OSError as exc:
        raise CorruptedFileError(f"Не удалось открыть файл «{path.name}»: {exc}") from exc

    if signature[:4] == _OLE2_SIGNATURE:
        return _load_xls(path)
    if signature[:2] == _ZIP_SIGNATURE:
        return _load_xlsx(path)
    raise NotExcelFileError(f"Файл «{path.name}» не является таблицей Excel (.xls или .xlsx).")
